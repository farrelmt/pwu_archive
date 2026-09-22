from datetime import date, timedelta
from io import BytesIO
from zipfile import ZipFile

from django.core import mail
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from pypdf import PdfReader, PdfWriter
from openpyxl import load_workbook

from accounts.models import ActivityLog, SystemUser
from .models import Disposisi, DisposisiLog, DisposisiRecipient


@override_settings(
    ALLOWED_HOSTS=["testserver"],
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    STORAGES={
        "default": {
            "BACKEND": "django.core.files.storage.InMemoryStorage",
        },
        "staticfiles": {
            "BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage",
        },
    },
)
class DisposisiSecurityTests(TestCase):
    def setUp(self):
        self.editor = SystemUser.objects.create_user(
            username="editor",
            password="password-for-tests",
            role="sekretaris",
        )
        self.admin = SystemUser.objects.create_user(
            username="admin-editor",
            password="password-for-tests",
            role="admin",
        )
        self.viewer = SystemUser.objects.create_user(
            username="viewer",
            password="password-for-tests",
            role="kadiv",
        )
        self.director = SystemUser.objects.create_user(
            username="director",
            password="password-for-tests",
            role="direktur_umum",
        )
        self.regular_director = SystemUser.objects.create_user(
            username="regular-director",
            password="password-for-tests",
            role="direktur",
        )
        self.wirajatim_user, _ = SystemUser.objects.update_or_create(
            username="wirajatim_kso",
            defaults={"role": "wirajatim_kso", "is_active": True},
        )
        self.main_director = SystemUser.objects.create_user(
            username="main-director",
            password="password-for-tests",
            role="direktur_utama",
        )
        self.accounting_head = SystemUser.objects.create_user(
            username="accounting-head",
            password="password-for-tests",
            role="kadiv_akuntansi",
        )
        self.finance_head = SystemUser.objects.create_user(
            username="finance-head",
            password="password-for-tests",
            role="kadiv_keuangan",
        )
        self.risk_head = SystemUser.objects.create_user(
            username="risk-head",
            password="password-for-tests",
            role="kadiv_risiko",
        )
        self.spi_observer = SystemUser.objects.create_user(
            username="spi-observer",
            password="password-for-tests",
            role="kadiv_spi",
        )
        self.disposisi = self.make_disposisi()

    def make_disposisi(self):
        return Disposisi.objects.create(
            tanggal_surat_diterima=date(2026, 7, 17),
            tanggal_surat=date(2026, 7, 16),
            nomor_surat="001/TEST",
            pengirim="Pengirim",
            lampiran="1",
            tujuan="DIR",
            tembusan="-",
            perihal="Test",
            tujuan_disposisi="Direktur",
            dokumen_surat_masuk=SimpleUploadedFile(
                "surat.pdf", b"%PDF-1.4 test", content_type="application/pdf"
            ),
        )

    def test_delete_rejects_get(self):
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:hapusdisposisi", args=[self.disposisi.pk])
        )

        self.assertEqual(response.status_code, 405)
        self.assertTrue(Disposisi.objects.filter(pk=self.disposisi.pk).exists())

    def test_spi_can_monitor_all_disposisi_but_cannot_modify_them(self):
        self.client.force_login(self.spi_observer)
        list_url = reverse("disposisi:disposisi")
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )

        list_response = self.client.get(list_url)
        self.assertEqual(list_response.status_code, 200)
        self.assertContains(list_response, self.disposisi.nomor_surat)
        self.assertNotContains(
            list_response,
            reverse("disposisi:tambahdisposisi"),
        )

        detail_response = self.client.get(detail_url)
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, self.disposisi.perihal)
        self.assertNotContains(
            detail_response,
            reverse("disposisi:editdisposisi", args=[self.disposisi.pk]),
        )

        protected_requests = (
            self.client.get(reverse("disposisi:tambahdisposisi")),
            self.client.get(
                reverse("disposisi:editdisposisi", args=[self.disposisi.pk])
            ),
            self.client.post(
                reverse("disposisi:hapusdisposisi", args=[self.disposisi.pk])
            ),
            self.client.get(
                reverse("disposisi:uploaddisposisi", args=[self.disposisi.pk])
            ),
        )
        for response in protected_requests:
            with self.subTest(path=response.request["PATH_INFO"]):
                self.assertEqual(response.status_code, 403)

        self.assertTrue(Disposisi.objects.filter(pk=self.disposisi.pk).exists())

    def test_create_redirects_to_new_disposisi_detail(self):
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:tambahdisposisi"),
            {
                "tanggal_surat_diterima": "2026-07-18",
                "tanggal_surat": "2026-07-17",
                "nomor_surat": "002/TEST",
                "pengirim": "Pengirim Baru",
                "lampiran": "1",
                "tujuan": "DIR",
                "tembusan": "-",
                "perihal": "Surat masuk baru",
                "dokumen_surat_masuk": SimpleUploadedFile(
                    "surat-baru.pdf",
                    b"%PDF-1.4 test",
                    content_type="application/pdf",
                ),
            },
        )

        created = Disposisi.objects.get(nomor_surat="002/TEST")
        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[created.pk]),
        )

    def test_create_warns_when_required_fields_are_missing(self):
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:tambahdisposisi"),
            {},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Data belum lengkap atau belum valid")
        for message in (
            "Tanggal surat diterima wajib diisi.",
            "Nomor surat wajib diisi.",
            "Pengirim wajib diisi.",
            "Perihal wajib diisi.",
            "Dokumen surat masuk wajib diunggah.",
        ):
            with self.subTest(message=message):
                self.assertContains(response, message)

    def test_create_missing_file_reminder_preserves_entered_values(self):
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:tambahdisposisi"),
            {
                "tanggal_surat_diterima": "2026-07-18",
                "tanggal_surat": "2026-07-17",
                "nomor_surat": "003/REMINDER",
                "pengirim": "Pengirim Pengingat",
                "lampiran": "1",
                "tujuan": "DIR",
                "tembusan": "-",
                "perihal": "Uji pengingat dokumen",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dokumen surat masuk wajib diunggah.")
        self.assertContains(response, 'value="003/REMINDER"')
        self.assertContains(response, "Uji pengingat dokumen")
        self.assertContains(response, 'id="formReminder"')
        self.assertContains(response, "Mohon isi atau unggah")
        self.assertFalse(
            Disposisi.objects.filter(nomor_surat="003/REMINDER").exists()
        )

    def test_non_editor_cannot_delete(self):
        self.client.force_login(self.viewer)

        response = self.client.post(
            reverse("disposisi:hapusdisposisi", args=[self.disposisi.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(Disposisi.objects.filter(pk=self.disposisi.pk).exists())

    def test_editor_can_delete_with_post(self):
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:hapusdisposisi", args=[self.disposisi.pk])
        )

        self.assertRedirects(response, reverse("disposisi:disposisi"))
        self.assertFalse(Disposisi.objects.filter(pk=self.disposisi.pk).exists())

    def test_upload_rejects_disallowed_file(self):
        self.client.force_login(self.editor)
        uploaded = SimpleUploadedFile(
            "payload.exe", b"not an image", content_type="application/octet-stream"
        )

        response = self.client.post(
            reverse("disposisi:uploadoffline", args=[self.disposisi.pk]),
            {
                "dokumen_disposisi": uploaded,
                "recipients": ["kadiv_akuntansi"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Format file tidak diizinkan")
        self.disposisi.refresh_from_db()
        self.assertFalse(self.disposisi.dokumen_disposisi)

    def test_upload_rejects_spoofed_image_content(self):
        self.client.force_login(self.editor)
        uploaded = SimpleUploadedFile(
            "fake.jpg",
            b"<html><script>alert(1)</script></html>",
            content_type="image/jpeg",
        )

        response = self.client.post(
            reverse("disposisi:uploadoffline", args=[self.disposisi.pk]),
            {
                "dokumen_disposisi": uploaded,
                "recipients": ["kadiv_akuntansi"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Isi file gambar tidak valid")

    def test_non_editor_cannot_upload(self):
        self.client.force_login(self.viewer)
        uploaded = SimpleUploadedFile(
            "disposisi.pdf", b"%PDF-1.4 test", content_type="application/pdf"
        )

        response = self.client.post(
            reverse("disposisi:uploadoffline", args=[self.disposisi.pk]),
            {"dokumen_disposisi": uploaded},
        )

        self.assertEqual(response.status_code, 403)

    def test_editor_detail_renders_tidy_created_action_toolbar(self):
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )

        for label in (
            "Kembali",
            "History",
            "Edit",
            "Dokumen Surat",
            "Preview",
            "Lanjutkan",
            "Metode Disposisi",
            "Belum Dipilih",
        ):
            with self.subTest(label=label):
                self.assertContains(response, label)
        self.assertContains(response, "fixed bottom-0")

    def test_detail_shows_offline_processing_method(self):
        self.disposisi.tipe_disposisi = "OFFLINE"
        self.disposisi.save(update_fields=["tipe_disposisi"])
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )

        self.assertContains(response, "Metode Disposisi")
        self.assertContains(response, "Offline")

    def test_history_uses_full_name_and_readable_username_fallback(self):
        self.director.first_name = "Direktur"
        self.director.last_name = "Utama"
        self.director.save(update_fields=["first_name", "last_name"])
        DisposisiLog.objects.create(
            disposisi=self.disposisi,
            user_log=self.director,
            action_log="SETUJUI_DISPOSISI",
        )
        DisposisiLog.objects.create(
            disposisi=self.disposisi,
            user_log=self.main_director,
            action_log="AJUKAN_DISPOSISI",
        )
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )

        self.assertContains(response, "Direktur Utama")
        self.assertContains(response, "Main Director")
        self.assertNotContains(response, ">director<")
        self.assertNotContains(response, ">main-director<")

    def test_forwarded_to_document_list_uses_wirajatim_kso_at_bottom(self):
        self.client.force_login(self.editor)
        preview_response = self.client.get(
            reverse("disposisi:previewdisposisi", args=[self.disposisi.pk])
        )

        self.assertNotContains(
            preview_response,
            '<span class="checkbox-text">Direktur</span>',
        )
        self.assertContains(preview_response, "Wirajatim KSO")
        preview_html = preview_response.content.decode()
        self.assertLess(
            preview_html.index("Kepala SPI"),
            preview_html.index("Wirajatim KSO"),
        )

        self.submit_online()
        self.client.force_login(self.director)
        online_response = self.client.get(
            reverse("disposisi:isionline", args=[self.disposisi.pk])
        )
        self.assertNotContains(online_response, "</span>Direktur</div>")
        self.assertContains(online_response, "</span>Wirajatim KSO</div>")
        online_html = online_response.content.decode()
        self.assertLess(
            online_html.index("Kepala SPI"),
            online_html.index("Wirajatim KSO"),
        )

    def submit_online(self, roles=None):
        self.client.force_login(self.editor)
        return self.client.post(
            reverse("disposisi:uploaddisposisi", args=[self.disposisi.pk]),
            {"metode": "ONLINE"},
        )

    def upload_offline(self, roles=None, user=None):
        self.client.force_login(user or self.editor)
        uploaded = SimpleUploadedFile(
            "disposisi.pdf",
            b"%PDF-1.4 offline test",
            content_type="application/pdf",
        )
        return self.client.post(
            reverse("disposisi:uploadoffline", args=[self.disposisi.pk]),
            {
                "dokumen_disposisi": uploaded,
                "recipients": roles or ["kadiv_akuntansi"],
            },
        )

    def test_method_page_links_to_dedicated_offline_upload(self):
        self.client.force_login(self.editor)

        page_response = self.client.get(
            reverse("disposisi:uploaddisposisi", args=[self.disposisi.pk])
        )

        offline_url = reverse(
            "disposisi:uploadoffline",
            args=[self.disposisi.pk],
        )
        self.assertContains(page_response, offline_url)
        self.assertContains(page_response, "Upload Offline")
        self.assertNotContains(page_response, "Bagikan Kepada")

        offline_response = self.client.get(offline_url)
        self.assertContains(offline_response, "Upload Disposisi Offline")
        self.assertContains(offline_response, "Bagikan Kepada")
        self.assertContains(offline_response, "Kepala Divisi Keuangan")
        self.assertContains(offline_response, 'value="wirajatim_kso"')
        self.assertContains(offline_response, "Wirajatim KSO")
        self.assertNotContains(offline_response, 'value="direktur"')

    def test_offline_can_be_shared_to_wirajatim_kso(self):
        response = self.upload_offline(["wirajatim_kso"])

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "SELESAI")
        recipient = self.disposisi.shared_recipients.get()
        self.assertEqual(recipient.role, "wirajatim_kso")
        self.assertEqual(recipient.get_role_display(), "Wirajatim KSO")

        preview_response = self.client.get(
            reverse("disposisi:previewdisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(
            preview_response,
            '<div class="checkbox">✓</div><span class="checkbox-text">Wirajatim KSO</span>',
            html=True,
        )

    def test_offline_upload_with_recipients_completes_progress(self):
        response = self.upload_offline(
            ["kadiv_akuntansi", "kadiv_keuangan"]
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.tipe_disposisi, "OFFLINE")
        self.assertEqual(self.disposisi.status_pengajuan, "SELESAI")
        self.assertSetEqual(
            set(self.disposisi.shared_recipients.values_list("role", flat=True)),
            {"kadiv_akuntansi", "kadiv_keuangan"},
        )

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        preview_response = self.client.get(
            reverse("disposisi:previewdisposisi", args=[self.disposisi.pk])
        )
        self.assertNotContains(detail_response, 'id="openShareModal"')
        self.assertContains(detail_response, "Kepala Divisi Akuntansi")
        self.assertContains(detail_response, "Kepala Divisi Keuangan")
        self.assertContains(preview_response, "✓", count=2)
        self.assertContains(
            preview_response,
            '<div class="checkbox">✓</div><span class="checkbox-text">Kepala Divisi Keuangan</span>',
            html=True,
        )
        self.assertTrue(
            self.disposisi.logs.filter(action_log="BAGI_DISPOSISI").exists()
        )
        self.assertTrue(self.disposisi.logs.filter(action_log="SELESAI").exists())
        expected_email_count = SystemUser.objects.filter(
            is_active=True,
            role__in=["kadiv_akuntansi", "kadiv_keuangan"],
        ).count()
        self.assertEqual(len(mail.outbox), expected_email_count)
        self.assertTrue(all(
            message.to == ["it.pwujatim@gmail.com"]
            for message in mail.outbox
        ))

    def test_uploaded_disposition_opens_preview_before_download(self):
        self.upload_offline()
        self.disposisi.refresh_from_db()
        self.client.force_login(self.editor)

        preview_url = reverse(
            "disposisi:preview_uploaded_disposisi",
            args=[self.disposisi.pk],
        )
        view_url = reverse(
            "disposisi:view_uploaded_disposisi",
            args=[self.disposisi.pk],
        )
        download_url = reverse(
            "disposisi:download_uploaded_disposisi",
            args=[self.disposisi.pk],
        )
        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        preview_response = self.client.get(preview_url)
        inline_response = self.client.get(view_url)
        download_response = self.client.get(download_url)

        self.assertContains(detail_response, preview_url)
        self.assertNotContains(
            detail_response,
            self.disposisi.dokumen_disposisi.url,
        )
        self.assertContains(preview_response, view_url)
        self.assertContains(preview_response, download_url)
        self.assertContains(preview_response, "Download")
        self.assertEqual(inline_response.status_code, 200)
        self.assertIn(
            "inline",
            inline_response["Content-Disposition"],
        )
        self.assertEqual(download_response.status_code, 200)
        self.assertIn(
            "attachment",
            download_response["Content-Disposition"],
        )

    def test_admin_role_cannot_upload_offline_disposition(self):
        response = self.upload_offline(["kadiv_aset"], user=self.admin)

        self.assertEqual(response.status_code, 403)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBUAT")
        self.assertFalse(
            self.disposisi.shared_recipients.filter(role="kadiv_aset").exists()
        )

    def approve_online(self, roles=None):
        self.submit_online(roles)
        self.client.force_login(self.director)
        return self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": "<p>Mohon ditindaklanjuti.</p>"},
        )

    def share_online(self, roles):
        self.approve_online()
        self.client.force_login(self.editor)
        return self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {
                "recipients": roles,
                "deadline": (
                    timezone.localdate() + timedelta(days=7)
                ).isoformat(),
            },
        )

    def test_editor_can_submit_online_request(self):
        self.disposisi.isi_disposisi = "<p>Isi lama.</p>"
        self.disposisi.save(update_fields=["isi_disposisi"])
        response = self.submit_online()

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.tipe_disposisi, "ONLINE")
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")
        self.assertEqual(self.disposisi.isi_disposisi, "")
        self.assertTrue(
            self.disposisi.logs.filter(action_log="AJUKAN_DISPOSISI").exists()
        )
        self.assertTrue(
            ActivityLog.objects.filter(
                actor=self.editor,
                category="DISPOSISI",
                action="AJUKAN_DISPOSISI",
                target_id=str(self.disposisi.pk),
            ).exists()
        )

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(detail_response, "Metode Disposisi")
        self.assertContains(detail_response, "Online")

    def test_pending_disposition_for_director_only_reaches_director(self):
        self.submit_online()
        isi_url = reverse("disposisi:isionline", args=[self.disposisi.pk])
        decision_url = reverse(
            "disposisi:decisiononline",
            args=[self.disposisi.pk],
        )

        self.client.force_login(self.main_director)
        wrong_monitor = self.client.get(reverse("homepage:monitor"))
        wrong_editor = self.client.get(isi_url)
        wrong_decision = self.client.post(decision_url, {"keputusan": "TOLAK"})

        self.assertNotContains(wrong_monitor, self.disposisi.nomor_surat)
        self.assertEqual(wrong_editor.status_code, 404)
        self.assertEqual(wrong_decision.status_code, 403)

        self.client.force_login(self.director)
        correct_monitor = self.client.get(reverse("homepage:monitor"))
        dashboard_response = self.client.get(reverse("homepage:dashboard"))
        self.assertContains(correct_monitor, self.disposisi.nomor_surat)
        self.assertContains(dashboard_response, 'id="notificationBadge"')
        self.assertContains(dashboard_response, "1 tugas perlu diproses")

    def test_pending_disposition_for_main_director_only_reaches_main_director(self):
        self.disposisi.tujuan = "DIRUT"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()

        self.client.force_login(self.director)
        regular_monitor = self.client.get(reverse("homepage:monitor"))
        self.assertNotContains(regular_monitor, self.disposisi.nomor_surat)

        self.client.force_login(self.main_director)
        main_monitor = self.client.get(reverse("homepage:inbox"))
        editor_response = self.client.get(
            reverse("disposisi:isionline", args=[self.disposisi.pk])
        )
        self.assertContains(main_monitor, self.disposisi.nomor_surat)
        self.assertEqual(editor_response.status_code, 200)
        self.assertContains(main_monitor, "Menunggu Direktur")

    def test_direksi_destination_reaches_both_directors(self):
        self.disposisi.tujuan = "DIREKSI"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()

        self.client.force_login(self.main_director)
        main_monitor = self.client.get(reverse("homepage:monitor"))

        self.client.force_login(self.director)
        director_monitor = self.client.get(reverse("homepage:monitor"))

        self.client.force_login(self.viewer)
        viewer_monitor = self.client.get(reverse("homepage:monitor"))

        self.assertContains(main_monitor, self.disposisi.nomor_surat)
        self.assertContains(director_monitor, self.disposisi.nomor_surat)
        self.assertNotContains(viewer_monitor, self.disposisi.nomor_surat)
        self.assertTrue(
            self.disposisi.can_be_approved_by(self.main_director)
        )
        self.assertTrue(self.disposisi.can_be_approved_by(self.director))

    def test_direksi_online_input_is_sequential(self):
        self.disposisi.tujuan = "DIREKSI"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()
        isi_url = reverse("disposisi:isionline", args=[self.disposisi.pk])

        self.client.force_login(self.regular_director)
        self.assertEqual(self.client.get(isi_url).status_code, 403)
        self.assertNotContains(
            self.client.get(reverse("homepage:inbox")),
            self.disposisi.nomor_surat,
        )

        self.client.force_login(self.main_director)
        self.assertContains(
            self.client.get(reverse("homepage:inbox")),
            self.disposisi.nomor_surat,
        )
        main_editor = self.client.get(isi_url)
        self.assertEqual(
            main_editor.context['editor_content'],
            '<div>Direktur Utama:</div><div>- </div>'
            '<div><br></div><div><br></div>',
        )
        self.client.post(
            isi_url,
            {"isi_disposisi": (
                '<p>Arahan pertama dari Dirut.</p>'
                '<svg data-signature-overlay="true" data-signature-layout="inline" '
                'data-signature-width="120" data-signature-height="50" '
                'data-signature-margin-left="300" data-signature-margin-top="8" '
                'data-signature-origin-x="300" data-signature-origin-y="250" '
                'viewBox="0 0 120 50">'
                '<path d="M 5 10 L 70 40"></path></svg>'
            )},
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")
        self.assertEqual(self.disposisi.isi_disposisi, "")
        self.assertIn("Arahan pertama", self.disposisi.isi_disposisi_dirut)
        self.assertIn(
            'data-signature-layout="positioned"',
            self.disposisi.isi_disposisi_dirut,
        )
        self.assertIn(
            'style="position: absolute; width: 120px; height: 50px; '
            'left: 300px; top: 250px"',
            self.disposisi.isi_disposisi_dirut,
        )

        self.client.force_login(self.regular_director)
        self.assertContains(
            self.client.get(reverse("homepage:inbox")),
            self.disposisi.nomor_surat,
        )
        director_editor = self.client.get(isi_url)
        self.assertEqual(
            director_editor.context['editor_content'],
            '<div>Direktur :</div><div>- </div>'
            '<div><br></div><div><br></div>',
        )
        self.client.post(
            isi_url,
            {"isi_disposisi": (
                '<p>Tindak lanjut dari Direktur.</p>'
                '<svg data-signature-overlay="true" viewBox="0 0 700 200">'
                '<path d="M 20 40 L 90 70"></path></svg>'
            )},
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertTrue(self.disposisi.direksi_input_complete)

        self.client.force_login(self.editor)
        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(detail_response, 'id="openShareModal"')

    def test_both_directors_get_one_combined_document_button_when_submitted(self):
        valid_letter = BytesIO()
        letter_writer = PdfWriter()
        letter_writer.add_blank_page(width=595, height=842)
        letter_writer.write(valid_letter)
        self.disposisi.dokumen_surat_masuk.save(
            "valid-letter.pdf",
            ContentFile(valid_letter.getvalue()),
            save=True,
        )
        self.disposisi.tujuan = "DIREKSI"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()

        combined_url = reverse(
            "disposisi:combined_director_document",
            args=[self.disposisi.pk],
        )
        preview_url = reverse(
            "disposisi:previewdisposisi",
            args=[self.disposisi.pk],
        )
        for user in (self.main_director, self.regular_director):
            with self.subTest(role=user.role):
                self.client.force_login(user)
                detail_response = self.client.get(
                    reverse(
                        "disposisi:detaildisposisi",
                        args=[self.disposisi.pk],
                    )
                )
                self.assertContains(detail_response, combined_url)
                self.assertContains(
                    detail_response,
                    "Lihat Isi Disposisi",
                    count=1,
                )
                self.assertNotContains(detail_response, "Dokumen Surat")
                self.assertNotContains(detail_response, preview_url)

                combined_response = self.client.get(combined_url)
                self.assertEqual(combined_response.status_code, 200)
                self.assertEqual(
                    combined_response["Content-Type"],
                    "application/pdf",
                )
                self.assertIn(
                    "inline",
                    combined_response["Content-Disposition"],
                )
                merged_pdf = PdfReader(BytesIO(combined_response.content))
                self.assertGreaterEqual(len(merged_pdf.pages), 2)
                self.assertIn(
                    "LEMBAR DISPOSISI",
                    merged_pdf.pages[0].extract_text(),
                )
                self.assertNotIn(
                    "LEMBAR DISPOSISI",
                    merged_pdf.pages[-1].extract_text() or "",
                )

    def test_non_director_cannot_open_combined_submitted_document(self):
        self.disposisi.tujuan = "DIREKSI"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse(
                "disposisi:combined_director_document",
                args=[self.disposisi.pk],
            )
        )

        self.assertEqual(response.status_code, 403)

    def test_both_directors_get_one_combined_preview_after_filling_disposition(self):
        valid_letter = BytesIO()
        letter_writer = PdfWriter()
        letter_writer.add_blank_page(width=595, height=842)
        letter_writer.write(valid_letter)
        self.disposisi.dokumen_surat_masuk.save(
            "completed-letter.pdf",
            ContentFile(valid_letter.getvalue()),
            save=True,
        )
        self.disposisi.tujuan = "DIREKSI"
        self.disposisi.save(update_fields=["tujuan"])
        self.submit_online()
        self.client.force_login(self.main_director)
        fill_response = self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {
                "isi_disposisi": (
                    '<p>Arahan Direktur Utama.</p>'
                    '<svg data-signature-overlay="true" viewBox="0 0 700 200">'
                    '<path d="M 10 40 L 80 70"></path></svg>'
                )
            },
        )
        self.assertEqual(fill_response.status_code, 302)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")
        self.assertIn(
            "Arahan Direktur Utama",
            self.disposisi.isi_disposisi_dirut,
        )
        self.assertEqual(self.disposisi.isi_disposisi_direktur, "")

        self.client.force_login(self.regular_director)
        director_editor = self.client.get(
            reverse("disposisi:isionline", args=[self.disposisi.pk])
        )
        self.assertContains(director_editor, "Arahan Direktur Utama")
        self.assertContains(director_editor, "Isi Disposisi Direktur")
        self.assertContains(
            director_editor,
            'aria-label="Isi disposisi Direktur Utama yang terkunci"',
        )
        editor_html = director_editor.content.decode()
        locked_start = editor_html.index(
            'aria-label="Isi disposisi Direktur Utama yang terkunci"'
        )
        locked_end = editor_html.index('</section>', locked_start)
        self.assertNotIn(
            'contenteditable="true"',
            editor_html[locked_start:locked_end],
        )
        self.assertIn('id="isiEditor"', editor_html)
        self.assertContains(
            director_editor,
            'data-signature-layout="positioned"',
        )
        fill_response = self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {
                "isi_disposisi": (
                    '<p>Tindak lanjut Direktur.</p>'
                    '<svg data-signature-overlay="true" viewBox="0 0 700 200">'
                    '<path d="M 20 50 L 90 80"></path></svg>'
                )
            },
        )
        self.assertEqual(fill_response.status_code, 302)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertIn(
            "Tindak lanjut Direktur",
            self.disposisi.isi_disposisi_direktur,
        )
        self.assertIn(
            "Arahan Direktur Utama",
            self.disposisi.isi_disposisi_dirut,
        )

        combined_url = reverse(
            "disposisi:combined_director_document",
            args=[self.disposisi.pk],
        )
        document_url = reverse(
            "disposisi:preview_document",
            args=[self.disposisi.pk, "surat-masuk"],
        )
        isi_url = reverse(
            "disposisi:isionline",
            args=[self.disposisi.pk],
        )
        disposition_preview_url = reverse(
            "disposisi:previewdisposisi",
            args=[self.disposisi.pk],
        )

        for user in (self.main_director, self.regular_director):
            with self.subTest(role=user.role):
                self.client.force_login(user)
                detail_response = self.client.get(
                    reverse(
                        "disposisi:detaildisposisi",
                        args=[self.disposisi.pk],
                    )
                )
                self.assertContains(detail_response, combined_url, count=1)
                self.assertEqual(
                    detail_response.context["combined_director_document_label"],
                    "Preview",
                )
                self.assertNotContains(detail_response, "Dokumen Surat")
                self.assertNotContains(detail_response, "Lihat Isi Disposisi")
                self.assertContains(detail_response, "Disetujui")
                self.assertNotContains(detail_response, "Disetujui Direktur")
                self.assertNotContains(detail_response, document_url)
                self.assertNotContains(detail_response, isi_url)
                self.assertNotContains(detail_response, disposition_preview_url)

                combined_response = self.client.get(combined_url)
                self.assertEqual(combined_response.status_code, 200)
                merged_pdf = PdfReader(BytesIO(combined_response.content))
                self.assertGreaterEqual(len(merged_pdf.pages), 2)
                self.assertIn(
                    "LEMBAR DISPOSISI",
                    merged_pdf.pages[0].extract_text(),
                )
                self.assertIn(
                    "Arahan Direktur Utama",
                    merged_pdf.pages[0].extract_text(),
                )
                self.assertIn(
                    "Tindak lanjut Direktur",
                    merged_pdf.pages[0].extract_text(),
                )
                self.assertNotIn(
                    "LEMBAR DISPOSISI",
                    merged_pdf.pages[-1].extract_text() or "",
                )

    def test_editor_can_cancel_pending_online_request(self):
        self.submit_online()
        self.disposisi.isi_disposisi = "<p>Draft yang harus dihapus.</p>"
        self.disposisi.save(update_fields=["isi_disposisi"])

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )

        self.assertContains(detail_response, "Menunggu Persetujuan")
        self.assertContains(detail_response, "Batalkan Disposisi Online")
        self.assertContains(
            detail_response,
            reverse("disposisi:cancelonline", args=[self.disposisi.pk]),
        )

        response = self.client.post(
            reverse("disposisi:cancelonline", args=[self.disposisi.pk])
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.tipe_disposisi, "BELUM")
        self.assertEqual(self.disposisi.status_pengajuan, "DIBUAT")
        self.assertEqual(self.disposisi.isi_disposisi, "")

    def test_director_sees_pending_request_and_can_approve(self):
        self.submit_online()
        self.client.force_login(self.director)

        monitor_response = self.client.get(reverse("homepage:monitor"))
        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        isi_url = reverse("disposisi:isionline", args=[self.disposisi.pk])
        editor_response = self.client.get(isi_url)

        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")

        send_response = self.client.post(
            isi_url,
            {
                "isi_disposisi": (
                    '<p onclick="alert(1)"><strong>Setujui</strong> pengadaan.</p>'
                    '<svg data-signature-overlay="true" viewBox="0 0 700 330" '
                    'onload="alert(1)"><path d="M 20 40 L 80 90" '
                    'stroke="red" onclick="alert(1)"></path>'
                    '<script>alert("svg-xss")</script></svg>'
                    '<script>alert("xss")</script>'
                )
            },
        )

        self.assertContains(monitor_response, self.disposisi.nomor_surat)
        self.assertContains(monitor_response, "table-fixed")
        self.assertContains(
            monitor_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.assertContains(detail_response, "Isi Disposisi")
        self.assertContains(detail_response, isi_url)
        self.assertNotContains(
            detail_response,
            reverse("disposisi:editdisposisi", args=[self.disposisi.pk]),
        )
        self.assertNotContains(detail_response, "Lanjutkan")
        self.assertContains(editor_response, 'contenteditable="true"')
        self.assertContains(editor_response, "editorToolbar")
        self.assertContains(editor_response, "Gambar langsung pada dokumen")
        self.assertContains(editor_response, "previewDocumentMount")
        self.assertNotContains(editor_response, "signatureModal")
        self.assertContains(editor_response, "width: 210mm")
        self.assertContains(editor_response, "height: 297mm")
        self.assertContains(editor_response, "max-height: 297mm")
        self.assertContains(editor_response, "aspect-ratio: 210 / 297")
        self.assertContains(
            editor_response,
            "const horizontalOffset = (",
        )
        self.assertContains(editor_response, "const verticalOffset =")
        self.assertContains(editor_response, "editor.clientHeight")
        self.assertContains(editor_response, "contentLimitWarning")
        self.assertContains(editor_response, "body > div:first-of-type > div.flex")
        self.assertContains(editor_response, "#isiDisposisiForm")
        self.assertNotContains(editor_response, 'id="drawingCancelRow"')
        self.assertContains(editor_response, "flex-wrap: nowrap")
        self.assertContains(editor_response, 'class="online-document-header"')
        self.assertContains(editor_response, 'class="destination-option"')
        self.assertContains(editor_response, 'class="isi-title"')
        self.assertContains(editor_response, "margin-top: 30px")
        self.assertNotContains(editor_response, 'class="isi-section border-')
        self.assertContains(editor_response, "Preview")
        self.assertContains(editor_response, "Kirim")
        self.assertRedirects(
            send_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.tipe_disposisi, "ONLINE")
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertIn("<strong>Setujui</strong>", self.disposisi.isi_disposisi)
        self.assertIn('data-signature-overlay="true"', self.disposisi.isi_disposisi)
        self.assertIn(
            'preserveAspectRatio="xMidYMax meet"',
            self.disposisi.isi_disposisi,
        )
        self.assertIn('stroke="#111827"', self.disposisi.isi_disposisi)
        self.assertNotIn("onclick", self.disposisi.isi_disposisi)
        self.assertNotIn('stroke="red"', self.disposisi.isi_disposisi)
        self.assertNotIn("onload", self.disposisi.isi_disposisi)
        self.assertNotIn("script", self.disposisi.isi_disposisi)
        self.assertTrue(
            self.disposisi.logs.filter(
                action_log="SETUJUI_DISPOSISI", user_log=self.director
            ).exists()
        )

    def test_director_can_reject_pending_request(self):
        self.submit_online()
        self.disposisi.isi_disposisi = "<p>Draft yang ditolak.</p>"
        self.disposisi.save(update_fields=["isi_disposisi"])
        self.client.force_login(self.director)

        self.client.post(
            reverse("disposisi:decisiononline", args=[self.disposisi.pk]),
            {"keputusan": "TOLAK", "alasan": "Data belum lengkap."},
        )

        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.tipe_disposisi, "BELUM")
        self.assertEqual(self.disposisi.status_pengajuan, "DIBUAT")
        self.assertEqual(self.disposisi.isi_disposisi, "")
        rejection = self.disposisi.logs.get(action_log="TOLAK_DISPOSISI")
        self.assertEqual(rejection.keterangan_log, "Data belum lengkap.")

    def test_recipients_are_selected_after_online_approval(self):
        self.submit_online()
        self.client.force_login(self.editor)

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertNotContains(detail_response, "Dibagikan Kepada")
        self.assertNotContains(detail_response, 'id="openShareModal"')
        self.assertFalse(self.disposisi.shared_recipients.exists())

        self.client.force_login(self.director)
        response = self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": "<p>Mohon ditindaklanjuti.</p>"},
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertFalse(self.disposisi.shared_recipients.exists())

        self.client.force_login(self.editor)
        approved_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(approved_detail, "Dibagikan Kepada")
        self.assertContains(approved_detail, "Belum ada penerima")
        self.assertContains(approved_detail, 'id="openShareModal"')
        self.assertContains(approved_detail, 'id="shareDeadline"')
        self.assertContains(approved_detail, 'type="date"')
        self.assertNotContains(approved_detail, 'id="openDeadlinePicker"')
        self.assertNotContains(approved_detail, 'value="direktur"')
        self.assertContains(approved_detail, 'value="wirajatim_kso"')

        response = self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {
                "recipients": ["kadiv_akuntansi", "kadiv_keuangan"],
                "deadline": (
                    timezone.localdate() + timedelta(days=7)
                ).isoformat(),
            },
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBAGIKAN")
        self.assertEqual(
            self.disposisi.deadline,
            timezone.localdate() + timedelta(days=7),
        )
        shared_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(shared_detail, "Preview")
        self.assertNotContains(shared_detail, "Lihat Isi Disposisi")
        self.assertSetEqual(
            set(self.disposisi.shared_recipients.values_list("role", flat=True)),
            {"kadiv_akuntansi", "kadiv_keuangan"},
        )
        share_log = self.disposisi.logs.get(action_log="BAGI_DISPOSISI")
        self.assertIn("Kepala Divisi Akuntansi", share_log.keterangan_log)
        self.assertIn("Kepala Divisi Keuangan", share_log.keterangan_log)
        expected_email_count = SystemUser.objects.filter(
            is_active=True,
            role__in=["kadiv_akuntansi", "kadiv_keuangan"],
        ).count()
        self.assertEqual(len(mail.outbox), expected_email_count)
        for notification in mail.outbox:
            self.assertTrue(
                notification.subject.startswith(
                    "NOTIFIKASI SISTEM ARSIP - "
                )
            )
            self.assertEqual(notification.to, ["it.pwujatim@gmail.com"])
            self.assertIn(
                f"Disposisi {self.disposisi.nomor_agenda}",
                notification.body,
            )
            self.assertIn(self.disposisi.nomor_surat, notification.body)
            self.assertIn(self.disposisi.perihal, notification.body)
            self.assertIn(
                self.disposisi.deadline.strftime("%d/%m/%Y"),
                notification.body,
            )
            self.assertIn(
                "https://archive.pwujatim.site/disposisi/",
                notification.body,
            )
        self.assertEqual(
            ActivityLog.objects.filter(
                action="DISPOSITION_EMAIL_SENT",
                target_id=str(self.disposisi.pk),
            ).count(),
            expected_email_count,
        )

        shared_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(shared_detail, 'id="sharedRecipientDescription"')
        self.assertContains(shared_detail, "Dibagikan Kepada")
        self.assertContains(shared_detail, "Kepala Divisi Akuntansi")
        self.assertContains(shared_detail, "Kepala Divisi Keuangan")
        self.assertContains(shared_detail, "Belum diterima", count=2)
        self.assertContains(shared_detail, "bg-red-100")
        self.assertContains(shared_detail, 'id="activityDeadline"')
        self.assertContains(
            shared_detail,
            self.disposisi.deadline.strftime("%d/%m/%Y"),
        )

    def test_incoming_mail_table_displays_colored_status_badge(self):
        self.client.force_login(self.editor)

        response = self.client.get(reverse("disposisi:disposisi"))

        self.assertContains(response, self.disposisi.get_status_pengajuan_display())
        self.assertContains(response, "rounded-full")
        self.assertContains(response, "bg-gray-200")
        self.assertContains(response, "text-gray-700")

    def test_only_selected_roles_see_shared_disposition_in_inbox(self):
        self.share_online(["kadiv_akuntansi", "kadiv_keuangan"])

        for selected_user in (self.accounting_head, self.finance_head):
            with self.subTest(role=selected_user.role):
                self.client.force_login(selected_user)
                response = self.client.get(reverse("homepage:inbox"))
                self.assertContains(response, self.disposisi.nomor_surat)
                self.assertContains(response, "Disposisi Dibagikan")
                self.assertContains(response, 'id="notificationBadge"')

        self.client.force_login(self.risk_head)
        risk_response = self.client.get(reverse("homepage:inbox"))
        self.assertNotContains(risk_response, self.disposisi.nomor_surat)
        self.assertNotContains(risk_response, 'id="notificationBadge"')

        self.client.force_login(self.editor)
        secretary_response = self.client.get(reverse("homepage:inbox"))
        self.assertNotContains(secretary_response, self.disposisi.nomor_surat)

    def test_recipient_cannot_select_online_recipients_after_approval(self):
        self.approve_online()
        self.client.force_login(self.accounting_head)

        response = self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {"recipients": ["kadiv_akuntansi"]},
        )

        self.assertEqual(response.status_code, 403)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertFalse(self.disposisi.shared_recipients.exists())

    def test_secretary_must_set_deadline_when_sharing_online(self):
        self.approve_online()
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {"recipients": ["kadiv_akuntansi"]},
            follow=True,
        )

        self.assertContains(response, "Deadline disposisi wajib dipilih.")
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertIsNone(self.disposisi.deadline)
        self.assertFalse(self.disposisi.shared_recipients.exists())

    def test_online_method_does_not_require_recipient_before_approval(self):
        self.client.force_login(self.editor)

        response = self.client.post(
            reverse("disposisi:uploaddisposisi", args=[self.disposisi.pk]),
            {"metode": "ONLINE"},
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")
        self.assertFalse(self.disposisi.shared_recipients.exists())

    def test_recipient_can_read_shared_online_content(self):
        self.share_online(["kadiv_akuntansi"])
        self.client.force_login(self.accounting_head)

        response = self.client.get(
            reverse("disposisi:isionline", args=[self.disposisi.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mohon ditindaklanjuti.")
        self.assertNotContains(response, 'role="textbox"')
        self.assertNotContains(response, ">Kirim<")

    def test_wirajatim_kso_can_receive_and_complete_online_disposisi(self):
        self.share_online(["wirajatim_kso"])
        self.client.force_login(self.wirajatim_user)
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )
        receive_url = reverse(
            "disposisi:receiveonline",
            args=[self.disposisi.pk],
        )
        complete_url = reverse(
            "disposisi:completeonline",
            args=[self.disposisi.pk],
        )

        detail_response = self.client.get(detail_url)
        self.assertContains(detail_response, "Wirajatim KSO")
        self.assertContains(detail_response, receive_url)
        self.assertContains(detail_response, "Terima")

        self.assertRedirects(self.client.post(receive_url), detail_url)
        self.assertRedirects(
            self.client.post(
                complete_url,
                {
                    "activity_description": "Tindak lanjut Wirajatim selesai.",
                    "follow_up_result": "Dokumen Wirajatim telah diproses.",
                },
            ),
            detail_url,
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")
        recipient = self.disposisi.shared_recipients.get(
            role="wirajatim_kso",
        )
        self.assertIsNotNone(recipient.received_at)
        self.assertIsNotNone(recipient.agreed_at)

    def test_director_recipients_are_informational_and_do_not_block_completion(self):
        self.share_online([
            "direktur_utama",
            "direktur_umum",
            "kadiv_akuntansi",
        ])
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )
        receive_url = reverse(
            "disposisi:receiveonline",
            args=[self.disposisi.pk],
        )
        complete_url = reverse(
            "disposisi:completeonline",
            args=[self.disposisi.pk],
        )

        for director in (
            self.main_director,
            self.director,
        ):
            with self.subTest(role=director.role):
                self.client.force_login(director)
                detail_response = self.client.get(detail_url)
                self.assertNotContains(detail_response, "Informasi saja")
                self.assertContains(detail_response, "Telah Dibagikan")
                self.assertNotContains(detail_response, receive_url)
                self.assertNotContains(detail_response, complete_url)

                self.assertRedirects(self.client.post(receive_url), detail_url)
                self.assertRedirects(
                    self.client.post(
                        complete_url,
                        {"activity_description": "Tidak diperlukan."},
                    ),
                    detail_url,
                )
                recipient = self.disposisi.shared_recipients.get(
                    role=director.role,
                )
                self.assertIsNone(recipient.received_at)
                self.assertIsNone(recipient.agreed_at)
                self.assertEqual(recipient.activity_description, "")

        self.client.force_login(self.accounting_head)
        self.client.post(receive_url)
        self.client.post(
            complete_url,
            {
                "activity_description": "Pemeriksaan akuntansi selesai.",
                "follow_up_result": "Pemeriksaan dinyatakan lengkap.",
            },
        )

        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")
        self.assertEqual(
            self.disposisi.shared_recipients.filter(
                agreed_at__isnull=False,
            ).count(),
            1,
        )

    def test_only_informational_directors_go_straight_to_verification(self):
        self.share_online([
            "direktur_utama",
            "direktur_umum",
        ])

        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")
        self.assertFalse(
            self.disposisi.shared_recipients.filter(
                received_at__isnull=False,
            ).exists()
        )

        self.client.force_login(self.editor)
        response = self.client.post(
            reverse("disposisi:approvecomplete", args=[self.disposisi.pk])
        )

        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "SELESAI")

    def test_spi_recipient_is_monitoring_only_and_requires_no_activity(self):
        self.share_online(["kadiv_spi"])
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")

        self.client.force_login(self.spi_observer)
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )
        detail_response = self.client.get(detail_url)
        self.assertContains(detail_response, "Telah Dibagikan")
        self.assertNotContains(
            detail_response,
            reverse("disposisi:receiveonline", args=[self.disposisi.pk]),
        )
        self.assertNotContains(
            detail_response,
            reverse("disposisi:completeonline", args=[self.disposisi.pk]),
        )

    def test_editing_shared_recipients_preserves_existing_activity_progress(self):
        self.share_online(["kadiv_akuntansi", "kadiv_keuangan"])
        self.disposisi.refresh_from_db()
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )

        self.client.force_login(self.accounting_head)
        self.client.post(
            reverse("disposisi:receiveonline", args=[self.disposisi.pk])
        )
        self.client.post(
            reverse("disposisi:completeonline", args=[self.disposisi.pk]),
            {
                "activity_description": "Aktivitas akuntansi sudah selesai.",
                "follow_up_result": "Dokumen akuntansi telah diarsipkan.",
            },
        )
        accounting_recipient = self.disposisi.shared_recipients.get(
            role="kadiv_akuntansi"
        )
        original_received_at = accounting_recipient.received_at
        original_agreed_at = accounting_recipient.agreed_at
        original_completed_by = accounting_recipient.completed_by

        self.client.force_login(self.editor)
        edit_detail = self.client.get(detail_url)
        self.assertContains(edit_detail, "Edit Penerima")
        self.assertContains(edit_detail, 'value="kadiv_akuntansi"')
        self.assertIn(
            "kadiv_akuntansi",
            edit_detail.context["selected_recipient_roles"],
        )
        response = self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {
                "recipients": [
                    "kadiv_akuntansi",
                    "kadiv_keuangan",
                    "kadiv_aset",
                ],
                "deadline": self.disposisi.deadline.isoformat(),
            },
        )

        self.assertRedirects(response, detail_url)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBAGIKAN")
        self.assertSetEqual(
            set(self.disposisi.shared_recipients.values_list("role", flat=True)),
            {"kadiv_akuntansi", "kadiv_keuangan", "kadiv_aset"},
        )
        accounting_recipient.refresh_from_db()
        self.assertEqual(accounting_recipient.received_at, original_received_at)
        self.assertEqual(accounting_recipient.agreed_at, original_agreed_at)
        self.assertEqual(
            accounting_recipient.activity_description,
            "Aktivitas akuntansi sudah selesai.",
        )
        self.assertEqual(accounting_recipient.completed_by, original_completed_by)
        added_recipient = self.disposisi.shared_recipients.get(role="kadiv_aset")
        self.assertIsNone(added_recipient.received_at)
        self.assertIsNone(added_recipient.agreed_at)
        self.assertEqual(added_recipient.activity_description, "")

    def test_editing_recipients_at_verification_returns_to_shared_progress(self):
        self.share_online(["kadiv_akuntansi"])
        detail_url = reverse(
            "disposisi:detaildisposisi",
            args=[self.disposisi.pk],
        )

        self.client.force_login(self.accounting_head)
        self.client.post(
            reverse("disposisi:receiveonline", args=[self.disposisi.pk])
        )
        self.client.post(
            reverse("disposisi:completeonline", args=[self.disposisi.pk]),
            {
                "activity_description": "Verifikasi akuntansi selesai.",
                "follow_up_result": "Data akuntansi telah sesuai.",
            },
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")
        accounting_recipient = self.disposisi.shared_recipients.get(
            role="kadiv_akuntansi"
        )
        original_agreed_at = accounting_recipient.agreed_at

        self.client.force_login(self.editor)
        verification_detail = self.client.get(detail_url)
        self.assertContains(verification_detail, "Edit Penerima")
        self.assertContains(verification_detail, "Edit Penerima Disposisi")
        response = self.client.post(
            reverse("disposisi:shareonline", args=[self.disposisi.pk]),
            {
                "recipients": ["kadiv_akuntansi", "kadiv_aset"],
                "deadline": self.disposisi.deadline.isoformat(),
            },
        )

        self.assertRedirects(response, detail_url)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBAGIKAN")
        accounting_recipient.refresh_from_db()
        self.assertEqual(accounting_recipient.agreed_at, original_agreed_at)
        self.assertEqual(
            accounting_recipient.activity_description,
            "Verifikasi akuntansi selesai.",
        )
        added_recipient = self.disposisi.shared_recipients.get(role="kadiv_aset")
        self.assertIsNone(added_recipient.received_at)
        self.assertIsNone(added_recipient.agreed_at)

    def test_disposition_finishes_after_all_recipients_submit_activity(self):
        self.share_online(["kadiv_akuntansi", "kadiv_keuangan"])
        self.client.force_login(self.accounting_head)
        receive_url = reverse(
            "disposisi:receiveonline",
            args=[self.disposisi.pk],
        )
        complete_url = reverse(
            "disposisi:completeonline",
            args=[self.disposisi.pk],
        )

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(detail_response, receive_url)
        self.assertContains(detail_response, "Terima")
        self.assertNotContains(detail_response, complete_url)

        receive_response = self.client.post(receive_url)

        self.assertRedirects(
            receive_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        accounting_recipient = self.disposisi.shared_recipients.get(
            role="kadiv_akuntansi"
        )
        self.assertIsNotNone(accounting_recipient.received_at)
        self.assertIsNone(accounting_recipient.agreed_at)
        acceptance_log = self.disposisi.logs.get(action_log="TERIMA_DISPOSISI")
        self.assertEqual(acceptance_log.user_log, self.accounting_head)
        self.assertIn("diterima", acceptance_log.keterangan_log)
        self.assertTrue(
            ActivityLog.objects.filter(
                actor=self.accounting_head,
                action="DISPOSITION_RECEIVED",
            ).exists()
        )

        monitor_response = self.client.get(reverse("homepage:monitor"))
        self.assertContains(monitor_response, self.disposisi.nomor_surat)

        received_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(received_detail, "Sudah diterima")
        self.assertContains(received_detail, "Belum diterima")
        self.assertContains(received_detail, "Receive Time")
        self.assertContains(
            received_detail,
            timezone.localtime(accounting_recipient.received_at).strftime(
                "%d/%m/%Y %H:%M"
            ),
        )
        self.assertContains(received_detail, "bg-yellow-100")
        self.assertContains(received_detail, "Done")
        self.assertContains(received_detail, complete_url)
        self.assertContains(received_detail, 'name="activity_description"')
        self.assertContains(received_detail, 'name="follow_up_result"')

        empty_activity_response = self.client.post(
            complete_url,
            {"activity_description": "   "},
            follow=True,
        )
        self.assertContains(
            empty_activity_response,
            "Aktivitas wajib diisi",
        )
        accounting_recipient.refresh_from_db()
        self.assertIsNone(accounting_recipient.agreed_at)

        empty_result_response = self.client.post(
            complete_url,
            {
                "activity_description": "Memeriksa dokumen.",
                "follow_up_result": "   ",
            },
            follow=True,
        )
        self.assertContains(
            empty_result_response,
            "Hasil tindak lanjut wajib diisi.",
        )
        accounting_recipient.refresh_from_db()
        self.assertIsNone(accounting_recipient.agreed_at)

        accounting_activity = "Memeriksa dan meneruskan dokumen ke tim akuntansi."
        accounting_result = "Dokumen telah diperiksa dan diteruskan."
        response = self.client.post(
            complete_url,
            {
                "activity_description": accounting_activity,
                "follow_up_result": accounting_result,
            },
        )
        self.assertRedirects(
            response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBAGIKAN")
        accounting_recipient.refresh_from_db()
        self.assertIsNotNone(accounting_recipient.agreed_at)
        self.assertEqual(
            accounting_recipient.activity_description,
            accounting_activity,
        )
        self.assertEqual(accounting_recipient.follow_up_result, accounting_result)
        self.assertEqual(
            accounting_recipient.completed_by,
            self.accounting_head,
        )
        activity_log = self.disposisi.logs.get(
            action_log="AKTIVITAS_PENERIMA",
            user_log=self.accounting_head,
        )
        self.assertIn(accounting_activity, activity_log.keterangan_log)
        self.assertFalse(self.disposisi.logs.filter(action_log="SELESAI").exists())

        activity_table_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(activity_table_detail, 'id="activityLogButton"')
        self.assertContains(
            activity_table_detail,
            'id="recipientActivityLogModal"',
        )
        self.assertContains(activity_table_detail, self.accounting_head.username)
        self.assertContains(activity_table_detail, accounting_activity)
        self.assertContains(activity_table_detail, accounting_result)
        self.assertContains(activity_table_detail, "Aktivitas yang Dilakukan")
        self.assertContains(activity_table_detail, "Hasil Tindak Lanjut")
        self.assertContains(activity_table_detail, "Edit Activity")
        self.assertContains(
            activity_table_detail,
            "waiting for other recipients",
        )
        history_actions = [
            log["action"]
            for daily_logs in activity_table_detail.context["grouped_logs"].values()
            for log in daily_logs
        ]
        self.assertNotIn("Aktivitas Penerima", history_actions)

        updated_accounting_activity = (
            "Memeriksa dokumen, meneruskan ke tim, dan mengarsipkan hasil."
        )
        edit_response = self.client.post(
            complete_url,
            {
                "activity_description": updated_accounting_activity,
                "follow_up_result": "Hasil tindak lanjut diperbarui.",
            },
        )
        self.assertRedirects(
            edit_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        accounting_recipient.refresh_from_db()
        self.assertEqual(
            accounting_recipient.activity_description,
            updated_accounting_activity,
        )
        self.assertTrue(
            ActivityLog.objects.filter(
                actor=self.accounting_head,
                action="RECIPIENT_ACTIVITY_UPDATED",
            ).exists()
        )

        self.client.force_login(self.finance_head)
        finance_monitor = self.client.get(reverse("homepage:monitor"))
        self.assertContains(finance_monitor, self.disposisi.nomor_surat)

        finance_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(finance_detail, receive_url)
        self.assertContains(finance_detail, "Terima")

        self.client.post(receive_url)
        finance_activity = "Melakukan verifikasi anggaran dan memberi catatan."

        final_response = self.client.post(
            complete_url,
            {
                "activity_description": finance_activity,
                "follow_up_result": "Anggaran telah diverifikasi.",
            },
        )
        self.assertRedirects(
            final_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")
        self.assertEqual(
            self.disposisi.shared_recipients.filter(
                agreed_at__isnull=False
            ).count(),
            2,
        )
        verification_log = self.disposisi.logs.get(
            action_log="AJUKAN_SELESAI"
        )
        self.assertEqual(verification_log.user_log, self.finance_head)
        self.assertIn("Menunggu persetujuan", verification_log.keterangan_log)
        self.assertFalse(self.disposisi.logs.filter(action_log="SELESAI").exists())

        completed_inbox = self.client.get(reverse("homepage:inbox"))
        self.assertNotContains(completed_inbox, self.disposisi.nomor_surat)

        pending_approval_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(pending_approval_detail, "Edit Activity")
        self.assertContains(pending_approval_detail, finance_activity)
        self.assertContains(
            pending_approval_detail,
            "waiting for Sekretaris approval",
        )

        edited_finance_activity = (
            "Melakukan verifikasi ulang anggaran dan menyimpan hasil final."
        )
        self.client.post(
            complete_url,
            {
                "activity_description": edited_finance_activity,
                "follow_up_result": "Verifikasi ulang selesai tanpa masalah.",
            },
        )
        finance_recipient = self.disposisi.shared_recipients.get(
            role="kadiv_keuangan"
        )
        self.assertEqual(
            finance_recipient.activity_description,
            edited_finance_activity,
        )

        self.client.force_login(self.editor)
        secretary_monitor = self.client.get(reverse("homepage:inbox"))
        self.assertContains(secretary_monitor, self.disposisi.nomor_surat)
        self.assertContains(
            secretary_monitor,
            "Menunggu Persetujuan Sekretaris",
        )

        approval_url = reverse(
            "disposisi:approvecomplete",
            args=[self.disposisi.pk],
        )
        secretary_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(secretary_detail, approval_url)
        self.assertContains(secretary_detail, "Approve Done")
        approval_response = self.client.post(approval_url)
        self.assertRedirects(
            approval_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "SELESAI")
        completion_log = self.disposisi.logs.get(action_log="SELESAI")
        self.assertEqual(completion_log.user_log, self.editor)
        self.assertIn("Sekretaris", completion_log.keterangan_log)

        self.client.force_login(self.finance_head)
        completed_detail = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        self.assertContains(completed_detail, "Selesai")
        self.assertContains(completed_detail, 'id="sharedRecipientDescription"')
        self.assertContains(completed_detail, "Kepala Divisi Akuntansi")
        self.assertContains(completed_detail, "Kepala Divisi Keuangan")
        self.assertContains(completed_detail, "(Selesai)", count=2)
        self.assertNotContains(completed_detail, "Belum diterima")
        self.assertNotContains(completed_detail, complete_url)

        locked_activity_response = self.client.post(
            complete_url,
            {"activity_description": "Tidak boleh berubah setelah approval."},
        )
        self.assertRedirects(
            locked_activity_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        finance_recipient.refresh_from_db()
        self.assertEqual(
            finance_recipient.activity_description,
            edited_finance_activity,
        )

    def test_recipient_must_receive_before_submitting_activity(self):
        self.share_online(["kadiv_akuntansi"])
        self.client.force_login(self.accounting_head)

        response = self.client.post(
            reverse("disposisi:completeonline", args=[self.disposisi.pk]),
            {"activity_description": "Mencoba langsung selesai."},
            follow=True,
        )

        self.assertContains(response, "Terima disposisi terlebih dahulu")
        recipient = self.disposisi.shared_recipients.get(
            role="kadiv_akuntansi"
        )
        self.assertIsNone(recipient.received_at)
        self.assertIsNone(recipient.agreed_at)
        self.assertEqual(recipient.activity_description, "")

    def test_non_secretary_cannot_approve_completion(self):
        self.share_online(["kadiv_akuntansi"])
        self.client.force_login(self.accounting_head)
        self.client.post(
            reverse("disposisi:receiveonline", args=[self.disposisi.pk])
        )
        self.client.post(
            reverse("disposisi:completeonline", args=[self.disposisi.pk]),
            {
                "activity_description": "Pekerjaan telah selesai.",
                "follow_up_result": "Dokumen telah ditindaklanjuti.",
            },
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")

        approval_url = reverse(
            "disposisi:approvecomplete",
            args=[self.disposisi.pk],
        )
        get_response = self.client.get(approval_url)
        post_response = self.client.post(approval_url)

        self.assertEqual(get_response.status_code, 405)
        self.assertEqual(post_response.status_code, 403)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "VERIFIKASI")

    def test_unselected_role_cannot_complete_shared_disposition(self):
        self.share_online(["kadiv_akuntansi"])
        self.client.force_login(self.finance_head)
        receive_url = reverse(
            "disposisi:receiveonline",
            args=[self.disposisi.pk],
        )
        complete_url = reverse(
            "disposisi:completeonline",
            args=[self.disposisi.pk],
        )

        receive_response = self.client.post(receive_url)
        get_response = self.client.get(complete_url)
        post_response = self.client.post(
            complete_url,
            {"activity_description": "Tidak berhak."},
        )

        self.assertEqual(receive_response.status_code, 403)
        self.assertEqual(get_response.status_code, 405)
        self.assertEqual(post_response.status_code, 403)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBAGIKAN")
        self.assertFalse(
            self.disposisi.logs.filter(action_log="TERIMA_DISPOSISI").exists()
        )
        self.assertFalse(self.disposisi.logs.filter(action_log="SELESAI").exists())

    def test_editor_can_open_approved_online_edit_without_file(self):
        self.submit_online()
        self.client.force_login(self.director)
        self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": "<p>Disetujui.</p>"},
        )
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:editdisposisi", args=[self.disposisi.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Belum ada file")
        self.assertNotContains(response, "data-dokumendisposisi=\" /media/")

        update_response = self.client.post(
            reverse("disposisi:editdisposisi", args=[self.disposisi.pk]),
            {
                "tanggal_surat_diterima": "2026-07-17",
                "tanggal_surat": "2026-07-16",
                "nomor_surat": "001/TEST",
                "pengirim": "Pengirim",
                "lampiran": "1",
                "tujuan": "DIR",
                "tembusan": "-",
                "perihal": "Test diperbarui",
            },
        )

        self.assertRedirects(
            update_response,
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk]),
        )
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIBUAT")
        self.assertEqual(self.disposisi.isi_disposisi, "")

    def test_non_director_cannot_decide_online_request(self):
        self.submit_online()

        response = self.client.post(
            reverse("disposisi:decisiononline", args=[self.disposisi.pk]),
            {"keputusan": "SETUJUI"},
        )

        self.assertEqual(response.status_code, 403)
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")

        editor_response = self.client.get(
            reverse("disposisi:isionline", args=[self.disposisi.pk])
        )
        self.assertEqual(editor_response.status_code, 403)

    def test_empty_online_content_does_not_approve(self):
        self.submit_online()
        self.client.force_login(self.director)

        response = self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": "<p><br></p>"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Isi disposisi wajib diisi")
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")

    def test_online_content_over_a4_limit_does_not_approve(self):
        self.submit_online()
        self.client.force_login(self.director)

        response = self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": f"<p>{'A' * 2401}</p>"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "melebihi batas satu halaman A4")
        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIAJUKAN")
        self.assertEqual(self.disposisi.isi_disposisi, "")

    def test_approved_online_content_is_read_only(self):
        self.submit_online()
        self.client.force_login(self.director)
        isi_url = reverse("disposisi:isionline", args=[self.disposisi.pk])
        self.client.post(isi_url, {"isi_disposisi": "<p>Mohon ditindaklanjuti.</p>"})
        self.client.force_login(self.editor)

        response = self.client.get(isi_url)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mohon ditindaklanjuti.")
        self.assertNotContains(response, 'role="textbox"')
        self.assertNotContains(response, ">Kirim<")

    def test_processed_request_cannot_be_decided_again(self):
        self.submit_online()
        self.client.force_login(self.director)
        decision_url = reverse(
            "disposisi:decisiononline", args=[self.disposisi.pk]
        )
        self.client.post(
            reverse("disposisi:isionline", args=[self.disposisi.pk]),
            {"isi_disposisi": "<p>Disetujui.</p>"},
        )

        self.client.post(decision_url, {"keputusan": "TOLAK"})

        self.disposisi.refresh_from_db()
        self.assertEqual(self.disposisi.status_pengajuan, "DIISI")
        self.assertFalse(
            self.disposisi.logs.filter(action_log="TOLAK_DISPOSISI").exists()
        )

    def test_unselected_user_cannot_list_or_open_disposition(self):
        self.client.force_login(self.risk_head)

        list_response = self.client.get(reverse("disposisi:disposisi"))
        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        preview_response = self.client.get(
            reverse("disposisi:previewdisposisi", args=[self.disposisi.pk])
        )

        self.assertEqual(list_response.status_code, 403)
        self.assertEqual(detail_response.status_code, 404)
        self.assertEqual(preview_response.status_code, 404)

    def test_selected_recipient_can_open_shared_document(self):
        self.disposisi.tipe_disposisi = "ONLINE"
        self.disposisi.status_pengajuan = "DIBAGIKAN"
        self.disposisi.save(
            update_fields=["tipe_disposisi", "status_pengajuan", "waktu_diedit"]
        )
        DisposisiRecipient.objects.create(
            disposisi=self.disposisi,
            role=self.accounting_head.role,
        )
        self.client.force_login(self.accounting_head)

        response = self.client.get(
            reverse(
                "disposisi:download_document",
                args=[self.disposisi.pk, "surat-masuk"],
            ),
            HTTP_X_FORWARDED_FOR="203.0.113.10",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["X-Accel-Redirect"].startswith("/protected-media/"))
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("private", response["Cache-Control"])
        self.assertIn("no-store", response["Cache-Control"])

    def test_incoming_document_opens_preview_before_download(self):
        self.client.force_login(self.editor)
        preview_url = reverse(
            "disposisi:preview_document",
            args=[self.disposisi.pk, "surat-masuk"],
        )
        view_url = reverse(
            "disposisi:view_document",
            args=[self.disposisi.pk, "surat-masuk"],
        )
        download_url = reverse(
            "disposisi:download_document",
            args=[self.disposisi.pk, "surat-masuk"],
        )

        detail_response = self.client.get(
            reverse("disposisi:detaildisposisi", args=[self.disposisi.pk])
        )
        preview_response = self.client.get(preview_url)
        inline_response = self.client.get(view_url)

        self.assertContains(detail_response, preview_url)
        self.assertContains(preview_response, view_url)
        self.assertContains(preview_response, download_url)
        self.assertContains(preview_response, "Download")
        self.assertEqual(inline_response.status_code, 200)
        self.assertNotIn("X-Accel-Redirect", inline_response)
        self.assertTrue(
            b"".join(inline_response.streaming_content).startswith(b"%PDF-1.4")
        )
        self.assertIn("inline", inline_response["Content-Disposition"])

    def test_unselected_user_cannot_download_document(self):
        self.client.force_login(self.risk_head)

        for route_name in (
            "preview_document",
            "view_document",
            "download_document",
        ):
            with self.subTest(route_name=route_name):
                response = self.client.get(
                    reverse(
                        f"disposisi:{route_name}",
                        args=[self.disposisi.pk, "surat-masuk"],
                    )
                )
                self.assertEqual(response.status_code, 404)

    def test_list_page_size_is_bounded(self):
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:disposisi"),
            {"limit": "999999999"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_limit"], "20")
        self.assertContains(response, 'id="exportSuratMasuk"')
        self.assertContains(response, reverse("disposisi:export_archive"))

    def test_list_filters_by_current_workflow_status_without_sender_filter(self):
        self.disposisi.status_pengajuan = "DIAJUKAN"
        self.disposisi.save(update_fields=["status_pengajuan"])
        completed = self.make_disposisi()
        completed.status_pengajuan = "SELESAI"
        completed.nomor_surat = "002/TEST"
        completed.save(update_fields=["status_pengajuan", "nomor_surat"])
        self.client.force_login(self.editor)

        response = self.client.get(
            reverse("disposisi:disposisi"),
            {"status": "diajukan", "pengirim": "Tidak Ada"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            list(response.context["page_obj"].object_list),
            [self.disposisi],
        )
        self.assertEqual(response.context["selected_status"], "DIAJUKAN")
        self.assertContains(response, "Disposisi Telah Diajukan")
        self.assertContains(response, "Menunggu Persetujuan Sekretaris")
        self.assertNotContains(response, 'id="filterPengirim"')

    def test_surat_masuk_archive_groups_files_and_excel_by_year(self):
        self.disposisi.tipe_disposisi = "ONLINE"
        self.disposisi.status_pengajuan = "DIISI"
        self.disposisi.isi_disposisi = "<p>Mohon <strong>ditindaklanjuti</strong>.</p>"
        self.disposisi.save(
            update_fields=["tipe_disposisi", "status_pengajuan", "isi_disposisi"]
        )
        DisposisiRecipient.objects.create(
            disposisi=self.disposisi,
            role="kadiv_keuangan",
        )
        excluded = self.make_disposisi()
        excluded.tanggal_surat_diterima = date(2025, 8, 1)
        excluded.tanggal_surat = date(2025, 7, 31)
        excluded.nomor_surat = "002/LAIN"
        excluded.perihal = "Tidak ikut ekspor"
        excluded.dokumen_disposisi = SimpleUploadedFile(
            "disposisi.png",
            b"not-a-real-image-but-valid-for-export",
            content_type="image/png",
        )
        excluded.save(
            update_fields=[
                "tanggal_surat_diterima",
                "tanggal_surat",
                "nomor_surat",
                "perihal",
                "dokumen_disposisi",
            ]
        )
        self.client.force_login(self.editor)

        response = self.client.get(reverse("disposisi:export_archive"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(".zip", response["Content-Disposition"])
        self.disposisi.refresh_from_db()
        excluded.refresh_from_db()
        agenda_2026 = self.disposisi.nomor_agenda.replace("/", "-")
        agenda_2025 = excluded.nomor_agenda.replace("/", "-")
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertIsNone(archive.testzip())
            archive_names = set(archive.namelist())
            self.assertIn("Surat Masuk/Daftar-Surat-Masuk.xlsx", archive_names)
            self.assertIn(
                f"Surat Masuk/2026/{agenda_2026}/Dokumen-Surat-Masuk.pdf",
                archive_names,
            )
            self.assertIn(
                f"Surat Masuk/2026/{agenda_2026}/Disposisi.pdf",
                archive_names,
            )
            self.assertIn(
                f"Surat Masuk/2025/{agenda_2025}/Dokumen-Surat-Masuk.pdf",
                archive_names,
            )
            self.assertIn(
                f"Surat Masuk/2025/{agenda_2025}/Disposisi.png",
                archive_names,
            )
            self.assertTrue(
                archive.read(
                    f"Surat Masuk/2026/{agenda_2026}/Disposisi.pdf"
                ).startswith(b"%PDF")
            )
            workbook = load_workbook(BytesIO(
                archive.read("Surat Masuk/Daftar-Surat-Masuk.xlsx")
            ))
        self.assertEqual(
            workbook.sheetnames,
            ["Surat Masuk 2026", "Surat Masuk 2025"],
        )
        detail_sheet = workbook["Surat Masuk 2026"]
        self.assertEqual(detail_sheet.max_row, 5)
        self.assertEqual(detail_sheet["E5"].value, "001/TEST")
        self.assertEqual(detail_sheet["L5"].value, "Disposisi Online")
        self.assertEqual(detail_sheet["N5"].value, "Mohon ditindaklanjuti.")
        self.assertEqual(detail_sheet["O5"].value, "Kepala Divisi Keuangan")
        self.assertEqual(detail_sheet["P5"].value, "PDF")
        self.assertEqual(
            detail_sheet["P5"].hyperlink.target,
            f"2026/{agenda_2026}/Dokumen-Surat-Masuk.pdf",
        )
        self.assertEqual(detail_sheet["Q5"].value, "PDF")
        self.assertEqual(
            detail_sheet["Q5"].hyperlink.target,
            f"2026/{agenda_2026}/Disposisi.pdf",
        )
        self.assertTrue(detail_sheet["T5"].hyperlink.target)
        self.assertEqual(detail_sheet["B5"].number_format, "dd/mm/yyyy")
        self.assertEqual(detail_sheet["A5"].alignment.horizontal, "center")
        self.assertEqual(detail_sheet["J5"].alignment.horizontal, "left")
        self.assertEqual(detail_sheet["N5"].alignment.horizontal, "left")
        self.assertEqual(detail_sheet.auto_filter.ref, "A4:T5")
        self.assertEqual(len(detail_sheet.tables), 0)
        for coordinate in ("A4", "A5", "J5", "N5", "T5"):
            with self.subTest(coordinate=coordinate):
                cell = detail_sheet[coordinate]
                self.assertEqual(cell.border.left.style, "thin")
                self.assertEqual(cell.border.right.style, "thin")
                self.assertEqual(cell.border.top.style, "thin")
                self.assertEqual(cell.border.bottom.style, "thin")
        older_sheet = workbook["Surat Masuk 2025"]
        self.assertEqual(older_sheet["E5"].value, "002/LAIN")
        self.assertEqual(older_sheet["Q5"].value, "IMG")
        self.assertEqual(
            older_sheet["Q5"].hyperlink.target,
            f"2025/{agenda_2025}/Disposisi.png",
        )

        filtered_response = self.client.get(
            reverse("disposisi:export_archive"),
            {"diterima_year": "2026"},
        )
        with ZipFile(BytesIO(filtered_response.content)) as archive:
            filtered_workbook = load_workbook(BytesIO(
                archive.read("Surat Masuk/Daftar-Surat-Masuk.xlsx")
            ))
        self.assertEqual(filtered_workbook.sheetnames, ["Surat Masuk 2026"])

    def test_surat_masuk_archive_export_requires_archive_list_access(self):
        self.client.force_login(self.viewer)

        response = self.client.get(reverse("disposisi:export_archive"))

        self.assertEqual(response.status_code, 403)

    def test_surat_masuk_archive_excel_escapes_formula_like_text(self):
        self.disposisi.pengirim = "=HYPERLINK(\"https://example.invalid\")"
        self.disposisi.save(update_fields=["pengirim"])
        self.client.force_login(self.editor)

        response = self.client.get(reverse("disposisi:export_archive"))
        with ZipFile(BytesIO(response.content)) as archive:
            workbook = load_workbook(
                BytesIO(archive.read("Surat Masuk/Daftar-Surat-Masuk.xlsx")),
                data_only=False,
            )

        self.assertEqual(
            workbook["Surat Masuk 2026"]["F5"].value,
            "'=HYPERLINK(\"https://example.invalid\")",
        )
        self.assertEqual(workbook["Surat Masuk 2026"]["F5"].data_type, "s")

    def test_surat_masuk_archive_excel_expands_rows_for_long_paragraphs(self):
        self.disposisi.isi_disposisi = (
            "<p>" + ("Uraian aktivitas yang panjang dan perlu dibaca. " * 80) + "</p>"
        )
        self.disposisi.dokumen_disposisi = SimpleUploadedFile(
            "disposisi.pdf",
            b"%PDF-1.4 test",
            content_type="application/pdf",
        )
        self.disposisi.save(
            update_fields=["isi_disposisi", "dokumen_disposisi"]
        )
        self.client.force_login(self.editor)

        response = self.client.get(reverse("disposisi:export_archive"))
        with ZipFile(BytesIO(response.content)) as archive:
            workbook = load_workbook(BytesIO(
                archive.read("Surat Masuk/Daftar-Surat-Masuk.xlsx")
            ))
        detail_sheet = workbook["Surat Masuk 2026"]

        self.assertGreater(detail_sheet.row_dimensions[5].height, 34)
        self.assertLessEqual(detail_sheet.row_dimensions[5].height, 409)
