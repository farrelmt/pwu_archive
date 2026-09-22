import mimetypes
import re
from urllib.parse import quote
import os
from io import BytesIO
from math import ceil
from zipfile import ZIP_DEFLATED, ZipFile

from django.shortcuts import render, get_object_or_404, redirect
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.http import FileResponse, Http404, JsonResponse, HttpResponse
from django.conf import settings
from django.template.loader import render_to_string
from django.db import transaction
from django.db.models import Q, Max
from django.urls import reverse
from django.utils.html import strip_tags
from .models import Disposisi, DisposisiLog, DisposisiRecipient
from .forms import (
    DisposisiForm,
    DisposisiUploadForm,
    OnlineDisposisiIsiForm,
    RecipientActivityForm,
    ShareDisposisiForm,
)
from .decorators import disposisi_director_required, disposisi_editor_required
from .access import visible_disposisi_for_user
from .notifications import send_disposition_shared_notifications
from datetime import datetime
from collections import defaultdict
from django.utils import timezone
from django.utils.timezone import localtime
from django.contrib import messages
from weasyprint import HTML
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError
from django.contrib.auth.decorators import login_required
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST
from django.utils.http import content_disposition_header
from accounts.audit import record_activity
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def notify_shared_recipients(request, disposisi, selected_roles):
    sent_count, failed_count = send_disposition_shared_notifications(
        request=request,
        disposisi=disposisi,
        recipient_roles=selected_roles,
    )
    if failed_count:
        messages.warning(
            request,
            f'{failed_count} email notifikasi gagal dikirim. '
            'Disposisi tetap berhasil dibagikan.',
        )
    return sent_count


def _filtered_disposisi_queryset(request):
    search = request.GET.get('search', '')
    surat_day = request.GET.get('surat_day')
    surat_month = request.GET.get('surat_month')
    surat_year = request.GET.get('surat_year')
    diterima_day = request.GET.get('diterima_day')
    diterima_month = request.GET.get('diterima_month')
    diterima_year = request.GET.get('diterima_year')

    tujuan = request.GET.get('tujuan')
    status = request.GET.get('status', '').strip().upper()
    status_choices = dict(Disposisi.STATUS_CHOICES)
    if status not in status_choices:
        status = ''

    #Get Disposisi Data
    sortDisposisi = request.GET.get('sort', 'tanggal_surat_diterima')
    orderDisposisi = request.GET.get('order', 'desc')

    ALLOWED_SORT = {
        'tsd': 'tanggal_surat_diterima',
        'ts': 'tanggal_surat',
        'na': 'nomor_agenda',
    }

    sortTable = ALLOWED_SORT.get(sortDisposisi, 'tanggal_surat_diterima')

    if orderDisposisi == 'asc':
        data = visible_disposisi_for_user(request.user).order_by(sortTable)
    else:
        data = visible_disposisi_for_user(request.user).order_by(f'-{sortTable}')

    SEARCH_FIELDS = [
        'tanggal_surat_diterima',
        'nomor_agenda',
        'tanggal_surat',
        'nomor_surat',
        'pengirim',
        'lampiran',
        'tembusan',
        'perihal',
        'tujuan',
    ]

    TUJUAN_MAP = {
        "direktur utama": "DIRUT",
        "dirut": "DIRUT",
        "direktur": "DIR",
        "dir": "DIR",
        "direksi": "DIREKSI",
    }

    STATUS_MAP = {
        "belum": "BELUM",
        "belum diajukan": "BELUM",
        "sudah": "SUDAH",
        "sudah diajukan": "SUDAH",
    }

    DATE_FORMATS = [
        "%Y-%m-%d",  # 2025-01-30
        "%d-%m-%Y",  # 30-01-2025
    ]

    if search:
        query = Q()
        for field in SEARCH_FIELDS:
            query |= Q(**{f"{field}__icontains": search})

        parsed_date = None
        for fmt in DATE_FORMATS:
            try:
                parsed_date = datetime.strptime(search, fmt)
                break
            except ValueError:
                continue

        # Search date with dd-mm-yy or yy-mm-dd
        if parsed_date:
            query |= Q(tanggal_surat_diterima=parsed_date)
            query |= Q(tanggal_surat=parsed_date)

        search_lower = search.lower()
        if search_lower in TUJUAN_MAP:
            query |= Q(tujuan=TUJUAN_MAP[search_lower])

        if search_lower in STATUS_MAP:
            query |= Q(status=STATUS_MAP[search_lower])

        data = data.filter(query)

    # search_lower = search.lower()
    # if search_lower in TUJUAN_MAP:
    #     data = data.filter(tujuan=TUJUAN_MAP[search_lower])

    #Use Filter
    if surat_year:
        data = data.filter(tanggal_surat__year=surat_year)

    if surat_month:
        data = data.filter(tanggal_surat__month=surat_month)

    if surat_day:
        data = data.filter(tanggal_surat__day=surat_day)

    if diterima_year:
        data = data.filter(tanggal_surat_diterima__year=diterima_year)

    if diterima_month:
        data = data.filter(tanggal_surat_diterima__month=diterima_month)

    if diterima_day:
        data = data.filter(tanggal_surat_diterima__day=diterima_day)

    if tujuan:
        data = data.filter(tujuan=tujuan)

    if status:
        data = data.filter(status_pengajuan=status)

    return data, search, status


@login_required
@never_cache
def list_disposisi(request):
    if not request.user.can_view_all_archive:
        raise PermissionDenied

    data, search, status = _filtered_disposisi_queryset(request)

    try:
        page_limit = int(request.GET.get('limit', 20))
    except ValueError:
        page_limit = 20
    if page_limit not in {20, 50, 100}:
        page_limit = 20

    try:
        page_number = int(request.GET.get('page', 1))
    except ValueError:
        page_number = 1

    paginator = Paginator(data, page_limit)
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'page_limit': str(page_limit),
        'search': search,
        'status_choices': Disposisi.STATUS_CHOICES,
        'selected_status': status,
    }

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        table_html = render_to_string('partials/disposisi_table.html', context, request=request)
        pagination_html = render_to_string('partials/disposisi_pagination.html', context, request=request)

        return JsonResponse({
            'table': table_html,
            'pagination': pagination_html,
        })

    return render(request, 'disposisi.html', context)


def _style_export_sheet(
    worksheet,
    header_row,
    widths,
    left_aligned_columns=(),
):
    dark_blue = '172554'
    light_blue = 'DBEAFE'
    thin_gray = Side(style='thin', color='CBD5E1')
    table_border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )
    left_aligned_columns = set(left_aligned_columns)

    worksheet.freeze_panes = f'A{header_row + 1}'
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 28

    for cell in worksheet[1]:
        cell.fill = PatternFill('solid', fgColor=dark_blue)
        cell.font = Font(color='FFFFFF', bold=True, size=14)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for cell in worksheet[header_row]:
        cell.fill = PatternFill('solid', fgColor=light_blue)
        cell.font = Font(color=dark_blue, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = table_border

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(widths),
    ):
        wrapped_line_count = 1
        for cell in row:
            cell.border = table_border
            cell.alignment = Alignment(
                horizontal=(
                    'left'
                    if cell.column in left_aligned_columns
                    else 'center'
                ),
                vertical='center',
                wrap_text=True,
            )
            if cell.column in left_aligned_columns:
                column_width = max(int(widths[cell.column - 1]), 1)
                text_lines = str(cell.value or '').splitlines() or ['']
                wrapped_line_count = max(
                    wrapped_line_count,
                    sum(
                        max(1, ceil(len(line) / column_width))
                        for line in text_lines
                    ),
                )
        worksheet.row_dimensions[row[0].row].height = min(
            409,
            max(34, wrapped_line_count * 15 + 8),
        )

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.row_dimensions[header_row].height = 34
    worksheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(widths))}{worksheet.max_row}'


def _excel_safe_text(value):
    text = str(value or '')
    if text.startswith(('=', '+', '-', '@')):
        return f"'{text}"
    return text


def _excel_document_label(document):
    extension = os.path.splitext(document.name)[1].lower()
    return 'PDF' if extension == '.pdf' else 'IMG'


_INVALID_ARCHIVE_COMPONENT = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _safe_archive_component(value, fallback):
    component = _INVALID_ARCHIVE_COMPONENT.sub('-', str(value or '')).strip(' .')
    return component or fallback


def _document_extension(document):
    return os.path.splitext(document.name)[1].lower() or '.bin'


def _stored_document_bytes(document):
    try:
        with document.open('rb') as source:
            return source.read()
    except (OSError, ValueError):
        return None


@login_required
@never_cache
@require_GET
def export_disposisi_archive(request):
    if not request.user.can_view_all_archive:
        raise PermissionDenied

    disposisi_list = list(
        _filtered_disposisi_queryset(request)[0].prefetch_related(
            'shared_recipients',
        )
    )
    disposisi_by_year = defaultdict(list)
    for disposisi in disposisi_list:
        disposisi_by_year[disposisi.tanggal_surat_diterima.year].append(disposisi)

    workbook = Workbook()
    workbook.remove(workbook.active)
    archive_entries = []
    used_archive_folders = set()
    generated_at = localtime(timezone.now()).replace(tzinfo=None)
    detail_headers = [
        'No', 'Tanggal Diterima', 'Nomor Agenda', 'Tanggal Surat',
        'Nomor Surat', 'Pengirim', 'Lampiran', 'Tujuan', 'Tembusan',
        'Perihal', 'Tujuan Disposisi', 'Metode Disposisi', 'Status',
        'Isi Disposisi', 'Dibagikan Kepada', 'Dokumen Surat Masuk',
        'Dokumen Disposisi', 'Waktu Dibuat', 'Waktu Diedit', 'Link Detail',
    ]
    yearly_groups = [
        (year, disposisi_by_year[year])
        for year in sorted(disposisi_by_year, reverse=True)
    ]
    if not yearly_groups:
        yearly_groups = [(None, [])]

    for year, yearly_disposisi in yearly_groups:
        sheet_suffix = str(year) if year is not None else ''
        worksheet = workbook.create_sheet(
            f'Surat Masuk {sheet_suffix}'.strip()
        )
        worksheet.append([
            f'DETAIL SURAT MASUK {sheet_suffix}'.strip()
        ])
        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(detail_headers),
        )
        worksheet.append(['Diekspor pada', generated_at])
        worksheet['B2'].number_format = 'dd/mm/yyyy hh:mm'
        worksheet.append([])
        worksheet.append(detail_headers)

        for number, disposisi in enumerate(yearly_disposisi, start=1):
            agenda_folder = _safe_archive_component(
                disposisi.nomor_agenda,
                f'Agenda-{disposisi.pk}',
            )
            relative_folder = f'{year}/{agenda_folder}'
            if relative_folder.casefold() in used_archive_folders:
                agenda_folder = f'{agenda_folder}-{disposisi.pk}'
                relative_folder = f'{year}/{agenda_folder}'
            used_archive_folders.add(relative_folder.casefold())

            incoming_document_path = None
            incoming_document_label = 'Tidak tersedia'
            incoming_document_bytes = _stored_document_bytes(
                disposisi.dokumen_surat_masuk
            )
            if incoming_document_bytes is not None:
                incoming_document_label = _excel_document_label(
                    disposisi.dokumen_surat_masuk
                )
                incoming_document_path = (
                    f'{relative_folder}/Dokumen-Surat-Masuk'
                    f'{_document_extension(disposisi.dokumen_surat_masuk)}'
                )
                archive_entries.append((
                    f'Surat Masuk/{incoming_document_path}',
                    incoming_document_bytes,
                ))

            disposition_document_path = None
            disposition_document_label = 'Tidak tersedia'
            if disposisi.dokumen_disposisi:
                disposition_document_bytes = _stored_document_bytes(
                    disposisi.dokumen_disposisi
                )
                if disposition_document_bytes is not None:
                    disposition_document_label = _excel_document_label(
                        disposisi.dokumen_disposisi
                    )
                    disposition_document_path = (
                        f'{relative_folder}/Disposisi'
                        f'{_document_extension(disposisi.dokumen_disposisi)}'
                    )
            else:
                disposition_document_bytes = _render_disposisi_pdf(
                    request,
                    disposisi,
                )
                disposition_document_label = 'PDF'
                disposition_document_path = f'{relative_folder}/Disposisi.pdf'
            if disposition_document_path and disposition_document_bytes is not None:
                archive_entries.append((
                    f'Surat Masuk/{disposition_document_path}',
                    disposition_document_bytes,
                ))

            recipient_names = ', '.join(
                recipient.get_role_display()
                for recipient in disposisi.shared_recipients.all()
            )
            if disposisi.tujuan == 'DIREKSI':
                isi_disposisi = '\n'.join(filter(None, [
                    'Direktur Utama: ' + ' '.join(strip_tags(
                        disposisi.isi_disposisi_dirut or ''
                    ).split()),
                    'Direktur: ' + ' '.join(strip_tags(
                        disposisi.isi_disposisi_direktur or ''
                    ).split()),
                ]))
            else:
                isi_disposisi = ' '.join(
                    strip_tags(disposisi.isi_disposisi or '').split()
                )
            detail_url = request.build_absolute_uri(
                reverse('disposisi:detaildisposisi', args=[disposisi.pk])
            )
            worksheet.append([
                number,
                disposisi.tanggal_surat_diterima,
                disposisi.nomor_agenda,
                disposisi.tanggal_surat,
                _excel_safe_text(disposisi.nomor_surat),
                _excel_safe_text(disposisi.pengirim),
                _excel_safe_text(disposisi.lampiran),
                disposisi.get_tujuan_display(),
                _excel_safe_text(disposisi.tembusan),
                _excel_safe_text(disposisi.perihal),
                _excel_safe_text(disposisi.tujuan_disposisi),
                disposisi.get_tipe_disposisi_display(),
                disposisi.get_status_pengajuan_display(),
                _excel_safe_text(isi_disposisi),
                recipient_names,
                incoming_document_label,
                disposition_document_label,
                localtime(disposisi.waktu_dibuat).replace(tzinfo=None),
                localtime(disposisi.waktu_diedit).replace(tzinfo=None),
                detail_url,
            ])
            incoming_document_cell = worksheet.cell(
                row=worksheet.max_row,
                column=16,
            )
            if incoming_document_path:
                incoming_document_cell.hyperlink = incoming_document_path
                incoming_document_cell.style = 'Hyperlink'
            disposition_document_cell = worksheet.cell(
                row=worksheet.max_row,
                column=17,
            )
            if disposition_document_path:
                disposition_document_cell.hyperlink = disposition_document_path
                disposition_document_cell.style = 'Hyperlink'
            detail_link_cell = worksheet.cell(
                row=worksheet.max_row,
                column=20,
            )
            detail_link_cell.hyperlink = detail_url
            detail_link_cell.style = 'Hyperlink'

        for row in worksheet.iter_rows(min_row=5):
            row[1].number_format = 'dd/mm/yyyy'
            row[3].number_format = 'dd/mm/yyyy'
            row[17].number_format = 'dd/mm/yyyy hh:mm'
            row[18].number_format = 'dd/mm/yyyy hh:mm'

        _style_export_sheet(
            worksheet,
            header_row=4,
            widths=[
                7, 17, 18, 17, 24, 24, 12, 20, 22, 45,
                24, 25, 28, 55, 35, 32, 32, 20, 20, 48,
            ],
            left_aligned_columns={10, 14},
        )
    workbook_output = BytesIO()
    workbook.save(workbook_output)

    archive_output = BytesIO()
    with ZipFile(archive_output, 'w', compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            'Surat Masuk/Daftar-Surat-Masuk.xlsx',
            workbook_output.getvalue(),
        )
        for archive_path, document_bytes in archive_entries:
            archive.writestr(archive_path, document_bytes)

    filename = f'surat-masuk-{generated_at:%Y%m%d-%H%M}.zip'
    response = HttpResponse(
        archive_output.getvalue(),
        content_type='application/zip',
    )
    response['Content-Disposition'] = content_disposition_header(True, filename)
    return response

def create_log(disposisi, user, action, desc=""):
    log = DisposisiLog.objects.create(
        disposisi= disposisi,
        user_log= user,
        action_log= action,
        keterangan_log  = desc
    )
    record_activity(
        actor=user,
        category='DISPOSISI',
        action=action,
        description=desc or log.get_action_log_display(),
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
        metadata={
            'status': disposisi.status_pengajuan,
            'method': disposisi.tipe_disposisi,
        },
    )


@login_required
@disposisi_editor_required
def tambah_disposisi(request):
    if request.method == "POST":
        form = DisposisiForm(request.POST, request.FILES)
        if form.is_valid():
            disposisi = form.save()

            create_log(disposisi, request.user, 'DIBUAT')
            return redirect('disposisi:detaildisposisi', pk=disposisi.pk)

    else:
        form = DisposisiForm()

    next_id_agenda =  Disposisi.objects.values('tanggal_surat_diterima').distinct().count() + 1

    return render(request, 'disposisi_tambah.html',{
        'form': form,
        'no_id_agenda': next_id_agenda,
    })

@login_required
@disposisi_editor_required
def update_disposisi(request, pk):
    disposisi = get_object_or_404(Disposisi, pk=pk)

    if request.method == "POST":
        form = DisposisiForm(request.POST, request.FILES, instance=disposisi)
        if form.is_valid():
            disposisi = form.save(commit=False)
            if disposisi.status_pengajuan == "DIISI" and disposisi.dokumen_disposisi:
                disposisi.dokumen_disposisi = None
            disposisi.isi_disposisi = ''
            disposisi.isi_disposisi_dirut = ''
            disposisi.isi_disposisi_direktur = ''
            disposisi.status_pengajuan = 'DIBUAT'
            disposisi.deadline = None
            disposisi.save()
            disposisi.shared_recipients.all().delete()

            create_log(disposisi, request.user, 'DIEDIT')
            return redirect('disposisi:detaildisposisi', pk=pk)
    else:
        form = DisposisiForm(instance=disposisi)

    return render(request, 'disposisi_edit.html', {
        'form': form,
        'disposisi': disposisi,
        'no_id_agenda': disposisi.id_agenda,
    })

@login_required
@disposisi_editor_required
@require_POST
def hapus_disposisi(request, pk):
    disposisi = get_object_or_404(Disposisi, pk=pk)
    files = [disposisi.dokumen_surat_masuk, disposisi.dokumen_disposisi]
    record_activity(
        request=request,
        category='DISPOSISI',
        action='HAPUS',
        description='Disposisi deleted.',
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
        metadata={
            'status': disposisi.status_pengajuan,
            'method': disposisi.tipe_disposisi,
        },
    )
    disposisi.delete()
    for document in files:
        if document:
            document.storage.delete(document.name)
    return redirect('disposisi:disposisi')

@login_required
@never_cache
def detail_disposisi(request, pk):
    disposisi = get_object_or_404(
        visible_disposisi_for_user(request.user).prefetch_related(
            'shared_recipients__completed_by',
        ),
        pk=pk,
    )
    share_recipient = disposisi.shared_recipients.filter(
        role=request.user.role,
    ).first()

    logs = disposisi.logs.select_related('user_log').exclude(
        action_log='AKTIVITAS_PENERIMA',
    )
    can_share_this_disposisi = (
        (
            disposisi.status_pengajuan == 'DIISI'
            or (
                disposisi.tipe_disposisi == 'ONLINE'
                and disposisi.status_pengajuan in {'DIBAGIKAN', 'VERIFIKASI'}
            )
        )
        and (
            (
                disposisi.tipe_disposisi == 'ONLINE'
                and request.user.can_share_disposisi
            )
            or (
                disposisi.tipe_disposisi == 'OFFLINE'
                and request.user.can_edit_disposisi
            )
        )
    )
    show_combined_director_document = (
        request.user.role in {'direktur_utama', 'direktur', 'direktur_umum'}
        and disposisi.tipe_disposisi == 'ONLINE'
        and disposisi.status_pengajuan in {'DIAJUKAN', 'DIISI'}
    )

    grouped_logs = defaultdict(list)

    HARI_ID = {
        'Monday': 'Senin',
        'Tuesday': 'Selasa',
        'Wednesday': 'Rabu',
        'Thursday': 'Kamis',
        'Friday': 'Jumat',
        'Saturday': 'Sabtu',
        'Sunday': 'Minggu',
    }

    for log in logs:
        local_time = localtime(log.waktu)
        hari_en = local_time.strftime('%A')
        hari_id = HARI_ID.get(hari_en, hari_en)
        date_key = f"{hari_id}, {local_time.strftime('%d/%m/%Y')}"
        user_name = log.user_log.get_full_name().strip()
        if not user_name:
            user_name = log.user_log.username.replace('_', ' ').replace('-', ' ').title()

        grouped_logs[date_key].append({
            'time': local_time.strftime('%H:%M'),
            'user': user_name,
            'action': log.get_action_log_display(),
            'desc': log.keterangan_log
        })


    return render(request, 'disposisi_detail.html', {
        'disposisi': disposisi,
        'grouped_logs': dict(grouped_logs),
        'share_role_choices': Disposisi.ONLINE_SHARE_ROLE_CHOICES,
        'can_approve_this_disposisi': disposisi.can_fill_online_disposition(
            request.user,
        ),
        'online_waiting_label': (
            'Menunggu Direktur'
            if (
                disposisi.tujuan == 'DIREKSI'
                and disposisi.isi_disposisi_dirut
                and not disposisi.isi_disposisi_direktur
            )
            else 'Menunggu Persetujuan'
        ),
        'is_share_recipient': share_recipient is not None,
        'is_actionable_share_recipient': (
            share_recipient is not None and share_recipient.requires_action
        ),
        'share_recipient': share_recipient,
        'recipient_activity_form': RecipientActivityForm(),
        'is_shared_online': (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan in {'DIBAGIKAN', 'VERIFIKASI'}
        ),
        'can_approve_completion': (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan == 'VERIFIKASI'
            and request.user.can_share_disposisi
        ),
        'can_share_this_disposisi': can_share_this_disposisi,
        'show_combined_director_document': show_combined_director_document,
        'combined_director_document_label': (
            'Preview'
            if disposisi.status_pengajuan == 'DIISI'
            else 'Lihat Isi Disposisi'
        ),
        'selected_recipient_roles': list(
            disposisi.shared_recipients.values_list('role', flat=True)
        ),
        'today_date': timezone.localdate().isoformat(),
        'is_deadline_overdue': (
            disposisi.deadline is not None
            and disposisi.deadline < timezone.localdate()
            and disposisi.status_pengajuan != 'SELESAI'
        ),
    })

@login_required
@never_cache
def preview_disposisi(request, pk):
    disposisi = get_object_or_404(
        visible_disposisi_for_user(request.user),
        pk=pk,
    )
    return render(request, 'disposisi_preview.html', {
        'disposisi': disposisi,
        'selected_recipient_roles': list(
            disposisi.shared_recipients.values_list('role', flat=True)
        ),
    })

@login_required
@disposisi_editor_required
def upload_disposisi(request, pk):
    disposisi = get_object_or_404(Disposisi, pk=pk)

    if request.method == "POST":
        if disposisi.status_pengajuan != 'DIBUAT':
            messages.error(request, "Disposisi ini sudah diproses.")
            return redirect('disposisi:detaildisposisi', pk=pk)

        metode = request.POST.get('metode')
        if metode == 'ONLINE':
            with transaction.atomic():
                disposisi = Disposisi.objects.select_for_update().get(pk=pk)
                if disposisi.status_pengajuan != 'DIBUAT':
                    messages.error(request, "Disposisi ini sudah diproses.")
                    return redirect('disposisi:detaildisposisi', pk=pk)
                disposisi.dokumen_disposisi = None
                disposisi.isi_disposisi = ''
                disposisi.isi_disposisi_dirut = ''
                disposisi.isi_disposisi_direktur = ''
                disposisi.deadline = None
                disposisi.tipe_disposisi = 'ONLINE'
                disposisi.status_pengajuan = 'DIAJUKAN'
                disposisi.save(update_fields=[
                    'dokumen_disposisi', 'isi_disposisi',
                    'isi_disposisi_dirut', 'isi_disposisi_direktur',
                    'deadline', 'tipe_disposisi',
                    'status_pengajuan', 'waktu_diedit'
                ])
                disposisi.shared_recipients.all().delete()
                create_log(
                    disposisi,
                    request.user,
                    'AJUKAN_DISPOSISI',
                    'Pengajuan disposisi online dikirim ke Direktur.',
                )
            messages.success(request, "Pengajuan online dikirim ke Direktur.")
            return redirect('disposisi:detaildisposisi', pk=pk)

        if metode == 'OFFLINE':
            return redirect('disposisi:uploadoffline', pk=pk)

        messages.error(request, "Pilih metode disposisi yang valid.")
        return redirect('disposisi:uploaddisposisi', pk=pk)

    return render(request, 'disposisi_upload.html', {
        'disposisi': disposisi,
    })


@login_required
@disposisi_editor_required
def upload_offline_disposisi(request, pk):
    disposisi = get_object_or_404(Disposisi, pk=pk)
    form = DisposisiUploadForm(
        request.POST or None,
        request.FILES or None,
        instance=disposisi,
    )
    share_form = ShareDisposisiForm(
        request.POST or None,
        choices=Disposisi.OFFLINE_SHARE_ROLE_CHOICES,
        include_deadline=False,
    )

    if request.method == "POST":
        if disposisi.status_pengajuan != 'DIBUAT':
            messages.error(request, "Disposisi ini sudah diproses.")
            return redirect('disposisi:detaildisposisi', pk=pk)

        if form.is_valid() and share_form.is_valid():
            selected_roles = share_form.cleaned_data['recipients']
            role_labels = dict(Disposisi.SHARE_ROLE_CHOICES)
            recipient_names = ', '.join(
                role_labels[role] for role in selected_roles
            )

            with transaction.atomic():
                locked_disposisi = Disposisi.objects.select_for_update().get(
                    pk=pk,
                )
                if locked_disposisi.status_pengajuan != 'DIBUAT':
                    messages.error(request, "Disposisi ini sudah diproses.")
                    return redirect('disposisi:detaildisposisi', pk=pk)

                locked_disposisi.dokumen_disposisi = (
                    form.cleaned_data['dokumen_disposisi']
                )
                locked_disposisi.tipe_disposisi = 'OFFLINE'
                locked_disposisi.status_pengajuan = 'SELESAI'
                locked_disposisi.save()
                locked_disposisi.shared_recipients.all().delete()
                DisposisiRecipient.objects.bulk_create([
                    DisposisiRecipient(disposisi=locked_disposisi, role=role)
                    for role in selected_roles
                ])
                create_log(
                    locked_disposisi,
                    request.user,
                    'UPLOAD_DISPOSISI',
                    'File disposisi offline berhasil diunggah.',
                )
                create_log(
                    locked_disposisi,
                    request.user,
                    'BAGI_DISPOSISI',
                    f'Disposisi dibagikan kepada: {recipient_names}.',
                )
                create_log(
                    locked_disposisi,
                    request.user,
                    'SELESAI',
                    'Disposisi offline selesai setelah file diunggah dan '
                    'penerima dipilih.',
                )

            notify_shared_recipients(request, locked_disposisi, selected_roles)
            messages.success(
                request,
                "File disposisi offline berhasil diunggah dan dibagikan.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)

        if share_form.errors.get('recipients'):
            messages.error(request, share_form.errors['recipients'][0])

    return render(request, 'disposisi_upload_offline.html', {
        'disposisi': disposisi,
        'form': form,
        'share_form': share_form,
    })


@login_required
@disposisi_editor_required
@require_POST
def cancel_online_disposisi(request, pk):
    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(), pk=pk
        )
        if not (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan == 'DIAJUKAN'
        ):
            messages.error(request, "Pengajuan online ini tidak dapat dibatalkan.")
            return redirect('disposisi:detaildisposisi', pk=pk)

        disposisi.tipe_disposisi = 'BELUM'
        disposisi.status_pengajuan = 'DIBUAT'
        disposisi.isi_disposisi = ''
        disposisi.isi_disposisi_dirut = ''
        disposisi.isi_disposisi_direktur = ''
        disposisi.deadline = None
        disposisi.save(update_fields=[
            'tipe_disposisi', 'status_pengajuan', 'isi_disposisi',
            'isi_disposisi_dirut', 'isi_disposisi_direktur', 'deadline',
            'waktu_diedit'
        ])
        disposisi.shared_recipients.all().delete()
        create_log(
            disposisi,
            request.user,
            'BATAL_PENGAJUAN',
            'Pengajuan disposisi online dibatalkan.',
        )

    messages.success(request, "Pengajuan online dibatalkan.")
    return redirect('disposisi:detaildisposisi', pk=pk)


@login_required
@disposisi_director_required
@require_POST
def decide_online_disposisi(request, pk):
    requested_disposisi = get_object_or_404(Disposisi, pk=pk)
    if not requested_disposisi.can_fill_online_disposition(request.user):
        raise PermissionDenied

    keputusan = request.POST.get('keputusan')
    if keputusan == 'SETUJUI':
        return redirect('disposisi:isionline', pk=pk)
    if keputusan != 'TOLAK':
        messages.error(request, "Keputusan tidak valid.")
        return redirect('homepage:monitor')

    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(), pk=pk
        )
        if not disposisi.can_fill_online_disposition(request.user):
            raise PermissionDenied
        if not (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan == 'DIAJUKAN'
        ):
            messages.error(request, "Pengajuan ini sudah diproses.")
            return redirect('homepage:monitor')

        alasan = request.POST.get('alasan', '').strip()
        disposisi.tipe_disposisi = 'BELUM'
        disposisi.status_pengajuan = 'DIBUAT'
        disposisi.isi_disposisi = ''
        disposisi.isi_disposisi_dirut = ''
        disposisi.isi_disposisi_direktur = ''
        disposisi.deadline = None
        action = 'TOLAK_DISPOSISI'
        description = alasan or 'Pengajuan online ditolak oleh Direktur.'
        success_message = "Pengajuan online ditolak."

        disposisi.save(update_fields=[
            'tipe_disposisi', 'status_pengajuan', 'isi_disposisi',
            'isi_disposisi_dirut', 'isi_disposisi_direktur', 'deadline',
            'waktu_diedit'
        ])
        disposisi.shared_recipients.all().delete()
        create_log(disposisi, request.user, action, description)

    messages.success(request, success_message)
    return redirect('homepage:monitor')


@login_required
@never_cache
def isi_online_disposisi(request, pk):
    disposisi = get_object_or_404(
        visible_disposisi_for_user(request.user),
        pk=pk,
    )
    is_pending = (
        disposisi.tipe_disposisi == 'ONLINE'
        and disposisi.status_pengajuan == 'DIAJUKAN'
    )
    is_approved = (
        disposisi.tipe_disposisi == 'ONLINE'
        and disposisi.status_pengajuan in {
            'DIISI', 'DIBAGIKAN', 'VERIFIKASI', 'SELESAI',
        }
    )

    if not (is_pending or is_approved):
        messages.error(request, "Disposisi online ini tidak tersedia untuk diisi.")
        return redirect('disposisi:detaildisposisi', pk=pk)

    can_approve_this = disposisi.can_fill_online_disposition(request.user)
    if is_pending and not can_approve_this:
        raise PermissionDenied

    read_only = is_approved
    if request.method == 'POST':
        if read_only or not can_approve_this:
            raise PermissionDenied

        stage = disposisi.online_input_stage_for(request.user)
        if stage is None:
            raise PermissionDenied
        stage_field, stage_label = stage
        form = OnlineDisposisiIsiForm(
            request.POST,
            max_layout_units=1100 if disposisi.tujuan == 'DIREKSI' else 2400,
            require_signature=disposisi.tujuan == 'DIREKSI',
        )
        if form.is_valid():
            with transaction.atomic():
                locked_disposisi = get_object_or_404(
                    Disposisi.objects.select_for_update(), pk=pk
                )
                if not (
                    locked_disposisi.tipe_disposisi == 'ONLINE'
                    and locked_disposisi.status_pengajuan == 'DIAJUKAN'
                ):
                    messages.error(request, "Pengajuan ini sudah diproses.")
                    return redirect('homepage:monitor')
                locked_stage = locked_disposisi.online_input_stage_for(
                    request.user,
                )
                if locked_stage is None or locked_stage[0] != stage_field:
                    raise PermissionDenied

                setattr(
                    locked_disposisi,
                    stage_field,
                    form.cleaned_data['isi_disposisi'],
                )
                if (
                    locked_disposisi.tujuan != 'DIREKSI'
                    or stage_field == 'isi_disposisi_direktur'
                ):
                    locked_disposisi.status_pengajuan = 'DIISI'
                locked_disposisi.save(update_fields=[
                    stage_field, 'status_pengajuan', 'waktu_diedit'
                ])
                create_log(
                    locked_disposisi,
                    request.user,
                    'SETUJUI_DISPOSISI',
                    f'Isi disposisi online dan tanda tangan {stage_label} '
                    'berhasil dikirim.',
                )
            if disposisi.tujuan == 'DIREKSI' and stage_field == 'isi_disposisi_dirut':
                messages.success(
                    request,
                    "Isi Direktur Utama berhasil dikirim. Menunggu Direktur.",
                )
            else:
                messages.success(
                    request,
                    "Isi disposisi berhasil dikirim dan disetujui.",
                )
            return redirect('disposisi:detaildisposisi', pk=pk)
    else:
        stage = disposisi.online_input_stage_for(request.user)
        stage_field, stage_label = stage or (None, 'Direktur')
        if read_only:
            if disposisi.tujuan == 'DIREKSI':
                editor_content = ''
            else:
                editor_content = disposisi.isi_disposisi
        else:
            editor_content = getattr(disposisi, stage_field, '')
            if disposisi.tujuan == 'DIREKSI' and not editor_content:
                heading = (
                    'Direktur Utama:'
                    if stage_field == 'isi_disposisi_dirut'
                    else 'Direktur :'
                )
                editor_content = (
                    f'<div>{heading}</div>'
                    '<div>- </div>'
                    '<div><br></div>'
                    '<div><br></div>'
                )
        form = OnlineDisposisiIsiForm(
            instance=disposisi,
            initial={'isi_disposisi': editor_content},
            max_layout_units=1100 if disposisi.tujuan == 'DIREKSI' else 2400,
            require_signature=disposisi.tujuan == 'DIREKSI',
        )

    return render(request, 'disposisi_isi_online.html', {
        'disposisi': disposisi,
        'form': form,
        'read_only': read_only,
        'editor_content': editor_content if request.method != 'POST' else request.POST.get('isi_disposisi', ''),
        'stage_label': stage_label,
        'is_dual_director': disposisi.tujuan == 'DIREKSI',
        'prior_director_content': disposisi.isi_disposisi_dirut,
        'selected_recipient_roles': list(
            disposisi.shared_recipients.values_list('role', flat=True)
        ),
    })


@login_required
@require_POST
def share_online_disposisi(request, pk):
    requested_disposisi = get_object_or_404(Disposisi, pk=pk)
    can_share = (
        request.user.can_share_disposisi
        if requested_disposisi.tipe_disposisi == 'ONLINE'
        else request.user.can_edit_disposisi
    )
    if not can_share:
        raise PermissionDenied

    form = ShareDisposisiForm(
        request.POST,
        existing_deadline=requested_disposisi.deadline,
    )
    if not form.is_valid():
        first_error = next(iter(form.errors.values()))[0]
        messages.error(request, first_error)
        return redirect('disposisi:detaildisposisi', pk=pk)

    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(),
            pk=pk,
        )
        can_update_recipients = (
            (
                disposisi.tipe_disposisi == 'ONLINE'
                and disposisi.status_pengajuan in {
                    'DIISI', 'DIBAGIKAN', 'VERIFIKASI'
                }
            )
            or (
                disposisi.tipe_disposisi == 'OFFLINE'
                and disposisi.status_pengajuan == 'DIISI'
            )
        )
        if not can_update_recipients:
            messages.error(request, "Disposisi ini belum siap dibagikan.")
            return redirect('disposisi:detaildisposisi', pk=pk)
        can_share_locked = (
            request.user.can_share_disposisi
            if disposisi.tipe_disposisi == 'ONLINE'
            else request.user.can_edit_disposisi
        )
        if not can_share_locked:
            raise PermissionDenied

        selected_roles = form.cleaned_data['recipients']
        deadline = form.cleaned_data['deadline']
        existing_roles = set(
            disposisi.shared_recipients.values_list('role', flat=True)
        )
        selected_role_set = set(selected_roles)
        added_roles = selected_role_set - existing_roles
        removed_roles = existing_roles - selected_role_set

        if removed_roles:
            disposisi.shared_recipients.filter(role__in=removed_roles).delete()
        DisposisiRecipient.objects.bulk_create([
            DisposisiRecipient(disposisi=disposisi, role=role)
            for role in added_roles
        ])
        is_offline = disposisi.tipe_disposisi == 'OFFLINE'
        has_pending_actionable_recipients = disposisi.shared_recipients.exclude(
            role__in=Disposisi.INFORMATIONAL_RECIPIENT_ROLES,
        ).filter(
            agreed_at__isnull=True,
        ).exists()
        has_actionable_recipients = disposisi.shared_recipients.exclude(
            role__in=Disposisi.INFORMATIONAL_RECIPIENT_ROLES,
        ).exists()
        recipients_were_updated = bool(existing_roles)
        if is_offline:
            disposisi.status_pengajuan = 'SELESAI'
        elif has_pending_actionable_recipients:
            disposisi.status_pengajuan = 'DIBAGIKAN'
        else:
            disposisi.status_pengajuan = 'VERIFIKASI'
        disposisi.deadline = deadline
        disposisi.save(update_fields=[
            'status_pengajuan', 'deadline', 'waktu_diedit'
        ])

        role_labels = dict(Disposisi.SHARE_ROLE_CHOICES)
        recipient_names = ', '.join(role_labels[role] for role in selected_roles)
        create_log(
            disposisi,
            request.user,
            'BAGI_DISPOSISI',
            (
                f'Penerima disposisi diperbarui menjadi: {recipient_names}. '
                f'Deadline: {deadline:%d/%m/%Y}.'
                if recipients_were_updated
                else f'Disposisi dibagikan kepada: {recipient_names}. '
                f'Deadline: {deadline:%d/%m/%Y}.'
            ),
        )
        if is_offline:
            create_log(
                disposisi,
                request.user,
                'SELESAI',
                'Disposisi offline selesai setelah file diunggah dan dibagikan.',
            )
        elif not has_actionable_recipients:
            create_log(
                disposisi,
                request.user,
                'AJUKAN_SELESAI',
                'Penerima hanya Direksi dan tidak memerlukan '
                'aktivitas. Menunggu persetujuan Sekretaris.',
            )

    notify_shared_recipients(request, disposisi, added_roles)
    if is_offline:
        messages.success(
            request,
            f"Disposisi offline berhasil dibagikan kepada {recipient_names} dan selesai.",
        )
    else:
        action_label = "Penerima disposisi berhasil diperbarui menjadi" if recipients_were_updated else "Disposisi berhasil dibagikan kepada"
        messages.success(request, f"{action_label} {recipient_names}.")
    return redirect('disposisi:detaildisposisi', pk=pk)


@login_required
@require_POST
def receive_shared_disposisi(request, pk):
    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(),
            pk=pk,
        )
        recipient = disposisi.shared_recipients.select_for_update().filter(
            role=request.user.role,
        ).first()
        if recipient is None:
            raise PermissionDenied
        if not recipient.requires_action:
            messages.info(
                request,
                "Penerima Direksi tidak perlu menerima disposisi.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)
        if not (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan == 'DIBAGIKAN'
        ):
            messages.error(request, "Disposisi ini tidak dapat diterima.")
            return redirect('disposisi:detaildisposisi', pk=pk)
        if recipient.agreed_at is not None:
            messages.info(request, "Disposisi ini sudah Anda selesaikan.")
            return redirect('disposisi:detaildisposisi', pk=pk)
        if recipient.received_at is not None:
            messages.info(request, "Disposisi ini sudah Anda terima.")
            return redirect('disposisi:detaildisposisi', pk=pk)

        recipient.received_at = timezone.now()
        recipient.save(update_fields=['received_at'])
        role_label = recipient.get_role_display()
        create_log(
            disposisi,
            request.user,
            'TERIMA_DISPOSISI',
            f'Disposisi diterima oleh {role_label}.',
        )

    record_activity(
        request=request,
        category='DISPOSISI',
        action='DISPOSITION_RECEIVED',
        description=f'Disposisi diterima oleh {role_label}.',
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
        metadata={'recipient_role': recipient.role},
    )
    messages.success(
        request,
        "Disposisi berhasil diterima. Selesaikan pekerjaan lalu isi aktivitas "
        "dan hasil tindak lanjut.",
    )
    return redirect('disposisi:detaildisposisi', pk=pk)


@login_required
@require_POST
def complete_shared_disposisi(request, pk):
    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(),
            pk=pk,
        )
        recipient = disposisi.shared_recipients.select_for_update().filter(
            role=request.user.role,
        ).first()
        if recipient is None:
            raise PermissionDenied
        if not recipient.requires_action:
            messages.info(
                request,
                "Penerima Direksi tidak perlu mengisi aktivitas.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)
        if not (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan in {'DIBAGIKAN', 'VERIFIKASI'}
        ):
            messages.error(request, "Disposisi ini tidak dapat diselesaikan.")
            return redirect('disposisi:detaildisposisi', pk=pk)
        if recipient.received_at is None:
            messages.error(
                request,
                "Terima disposisi terlebih dahulu sebelum menyelesaikannya.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)

        is_edit = recipient.agreed_at is not None
        form = RecipientActivityForm(request.POST)
        if not form.is_valid():
            first_error = next(iter(form.errors.values()))[0]
            messages.error(request, first_error)
            return redirect('disposisi:detaildisposisi', pk=pk)

        activity_description = form.cleaned_data['activity_description']
        follow_up_result = form.cleaned_data['follow_up_result']
        recipient.activity_description = activity_description
        recipient.follow_up_result = follow_up_result
        recipient.completed_by = request.user
        recipient.agreed_at = timezone.now()
        recipient.save(update_fields=[
            'activity_description',
            'follow_up_result',
            'completed_by',
            'agreed_at',
        ])
        role_label = recipient.get_role_display()
        create_log(
            disposisi,
            request.user,
            'AKTIVITAS_PENERIMA',
            f'{role_label} — Aktivitas: {activity_description} — '
            f'Hasil: {follow_up_result}',
        )

        is_complete = not disposisi.shared_recipients.exclude(
            role__in=Disposisi.INFORMATIONAL_RECIPIENT_ROLES,
        ).filter(
            agreed_at__isnull=True,
        ).exists()
        if is_complete and disposisi.status_pengajuan == 'DIBAGIKAN':
            disposisi.status_pengajuan = 'VERIFIKASI'
            disposisi.save(update_fields=['status_pengajuan', 'waktu_diedit'])
            create_log(
                disposisi,
                request.user,
                'AJUKAN_SELESAI',
                'Seluruh penerima telah menyelesaikan disposisi dan '
                'mengisi aktivitas serta hasil tindak lanjut. '
                'Menunggu persetujuan Sekretaris.',
            )

    record_activity(
        request=request,
        category='DISPOSISI',
        action=(
            'RECIPIENT_ACTIVITY_UPDATED'
            if is_edit
            else 'RECIPIENT_ACTIVITY_SUBMITTED'
        ),
        description=(
            f'{role_label} — Aktivitas: {activity_description} — '
            f'Hasil: {follow_up_result}'
        ),
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
        metadata={
            'recipient_role': recipient.role,
            'follow_up_result': follow_up_result,
            'all_recipients_complete': is_complete,
        },
    )
    if is_edit:
        messages.success(request, "Aktivitas berhasil diperbarui.")
    elif is_complete:
        messages.success(
            request,
            "Semua penerima telah mengisi aktivitas. "
            "Menunggu persetujuan Sekretaris.",
        )
    else:
        messages.success(
            request,
            "Aktivitas berhasil dikirim. Menunggu penerima lainnya.",
        )
    return redirect('disposisi:detaildisposisi', pk=pk)


@login_required
@require_POST
def approve_completed_disposisi(request, pk):
    if not request.user.can_share_disposisi:
        raise PermissionDenied

    with transaction.atomic():
        disposisi = get_object_or_404(
            Disposisi.objects.select_for_update(),
            pk=pk,
        )
        if not (
            disposisi.tipe_disposisi == 'ONLINE'
            and disposisi.status_pengajuan == 'VERIFIKASI'
        ):
            messages.error(
                request,
                "Disposisi ini belum siap disetujui sebagai selesai.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)
        if disposisi.shared_recipients.exclude(
            role__in=Disposisi.INFORMATIONAL_RECIPIENT_ROLES,
        ).filter(
            agreed_at__isnull=True,
        ).exists():
            messages.error(
                request,
                "Masih ada penerima yang belum mengisi aktivitas.",
            )
            return redirect('disposisi:detaildisposisi', pk=pk)

        disposisi.status_pengajuan = 'SELESAI'
        disposisi.save(update_fields=['status_pengajuan', 'waktu_diedit'])
        create_log(
            disposisi,
            request.user,
            'SELESAI',
            'Sekretaris menyetujui penyelesaian disposisi setelah '
            'memeriksa aktivitas seluruh penerima.',
        )

    record_activity(
        request=request,
        category='DISPOSISI',
        action='COMPLETION_APPROVED',
        description='Sekretaris menyetujui disposisi sebagai selesai.',
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
    )
    messages.success(request, "Penyelesaian disposisi telah disetujui.")
    return redirect('disposisi:detaildisposisi', pk=pk)


@login_required
@disposisi_editor_required
@require_POST
def edit_file_disposisi(request, pk):
    disposisi = get_object_or_404(Disposisi, pk=pk)

    form = DisposisiUploadForm(request.POST, request.FILES, instance=disposisi)
    if form.is_valid():
        disposisi = form.save()
        create_log(disposisi, request.user, 'DIEDIT')
        return redirect('disposisi:detaildisposisi', pk=pk)
    messages.error(request, "File tidak valid.")
    return redirect('disposisi:detaildisposisi', pk=pk)

@login_required
@never_cache
def download_disposisi_pdf(request, pk):
    disposisi = get_object_or_404(
        visible_disposisi_for_user(request.user),
        pk=pk,
    )
    record_activity(
        request=request,
        category='DISPOSISI',
        action='DOWNLOAD_PDF',
        description='Disposition PDF generated and downloaded.',
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
    )

    html_string = render_to_string(
        'disposisi_pdf.html',
        {
            'disposisi': disposisi,
            'selected_recipient_roles': list(
                disposisi.shared_recipients.values_list('role', flat=True)
            ),
        }
    )

    response = HttpResponse(content_type='application/pdf')
    filename = f'Disposisi-{disposisi.nomor_agenda}.pdf'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    HTML(
        string=html_string,
        base_url=request.build_absolute_uri('/')
    ).write_pdf(response)

    return response


def _render_disposisi_pdf(request, disposisi):
    html_string = render_to_string(
        'disposisi_pdf.html',
        {
            'disposisi': disposisi,
            'selected_recipient_roles': list(
                disposisi.shared_recipients.values_list('role', flat=True)
            ),
        },
    )
    return HTML(
        string=html_string,
        base_url=request.build_absolute_uri('/'),
    ).write_pdf()


def _incoming_letter_pdf(document):
    filename = document.name.rsplit('/', 1)[-1]
    content_type = mimetypes.guess_type(filename)[0]
    with document.open('rb') as source:
        document_bytes = source.read()
    if content_type == 'application/pdf' or filename.lower().endswith('.pdf'):
        return document_bytes

    with Image.open(BytesIO(document_bytes)) as image:
        if image.mode not in {'RGB', 'L'}:
            image = image.convert('RGB')
        output = BytesIO()
        image.save(output, format='PDF', resolution=150.0)
        return output.getvalue()


@login_required
@require_GET
@never_cache
def combined_director_document(request, pk):
    if request.user.role not in {'direktur_utama', 'direktur', 'direktur_umum'}:
        raise PermissionDenied
    disposisi = get_object_or_404(
        visible_disposisi_for_user(request.user),
        pk=pk,
        tipe_disposisi='ONLINE',
        status_pengajuan__in={'DIAJUKAN', 'DIISI'},
    )
    if not disposisi.dokumen_surat_masuk:
        raise Http404

    writer = PdfWriter()
    try:
        letter_reader = PdfReader(BytesIO(
            _incoming_letter_pdf(disposisi.dokumen_surat_masuk)
        ))
        preview_reader = PdfReader(BytesIO(
            _render_disposisi_pdf(request, disposisi)
        ))
        for page in preview_reader.pages:
            writer.add_page(page)
        for page in letter_reader.pages:
            writer.add_page(page)
    except (OSError, ValueError, PdfReadError) as exc:
        raise Http404('Dokumen PDF tidak dapat diproses.') from exc

    output = BytesIO()
    writer.write(output)
    response = HttpResponse(output.getvalue(), content_type='application/pdf')
    filename = f'Dokumen-Surat-{disposisi.nomor_agenda}.pdf'
    response['Content-Disposition'] = content_disposition_header(
        as_attachment=False,
        filename=filename,
    )
    response['Cache-Control'] = 'private, no-store, max-age=0'
    response['X-Content-Type-Options'] = 'nosniff'
    record_activity(
        request=request,
        category='DISPOSISI',
        action='VIEW_COMBINED_PDF',
        description='Incoming letter and disposition preview opened as one PDF.',
        target_type='disposisi.Disposisi',
        target_id=disposisi.pk,
        target_label=disposisi.nomor_agenda or disposisi.nomor_surat,
    )
    return response


def _document_or_404(user, pk, kind):
    field_name = {
        "surat-masuk": "dokumen_surat_masuk",
        "disposisi": "dokumen_disposisi",
    }.get(kind)
    if field_name is None:
        raise Http404

    disposisi = get_object_or_404(
        visible_disposisi_for_user(user),
        pk=pk,
    )
    document = getattr(disposisi, field_name)
    if not document:
        raise Http404
    return disposisi, document


def _protected_document_response(request, document, *, as_attachment):
    """Serve an authorized document through Nginx or directly in local use."""
    filename = document.name.rsplit("/", 1)[-1]
    content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    if request.META.get("HTTP_X_FORWARDED_FOR"):
        response = HttpResponse(content_type=content_type)
        response["Content-Disposition"] = content_disposition_header(
            as_attachment=as_attachment,
            filename=filename,
        )
        response["X-Accel-Redirect"] = (
            f"/protected-media/{quote(document.name, safe='/')}"
        )
    else:
        try:
            response = FileResponse(
                document.open("rb"),
                as_attachment=as_attachment,
                filename=filename,
                content_type=content_type,
            )
        except (OSError, ValueError) as exc:
            raise Http404("Dokumen tidak dapat dibuka.") from exc
    response["Cache-Control"] = "private, no-store, max-age=0"
    response["X-Content-Type-Options"] = "nosniff"
    return response


@login_required
@require_GET
@never_cache
def preview_document(request, pk, kind):
    disposisi, document = _document_or_404(request.user, pk, kind)
    return render(request, 'disposisi_document_file_preview.html', {
        'disposisi': disposisi,
        'kind': kind,
        'filename': os.path.basename(document.name),
        'document_label': (
            'Dokumen Surat' if kind == 'surat-masuk' else 'Dokumen Disposisi'
        ),
    })


@login_required
@require_GET
@never_cache
def view_document(request, pk, kind):
    _, document = _document_or_404(request.user, pk, kind)
    return _protected_document_response(
        request,
        document,
        as_attachment=False,
    )


@login_required
@require_GET
@never_cache
def download_document(request, pk, kind):
    """Authorize a document and return it as a download."""
    _, document = _document_or_404(request.user, pk, kind)
    return _protected_document_response(
        request,
        document,
        as_attachment=True,
    )


def _uploaded_disposisi_or_404(user, pk):
    disposisi = get_object_or_404(
        visible_disposisi_for_user(user),
        pk=pk,
    )
    if not disposisi.dokumen_disposisi:
        raise Http404
    return disposisi


@login_required
@require_GET
@never_cache
def preview_uploaded_disposisi(request, pk):
    disposisi = _uploaded_disposisi_or_404(request.user, pk)
    return render(request, 'disposisi_document_preview.html', {
        'disposisi': disposisi,
        'filename': os.path.basename(disposisi.dokumen_disposisi.name),
    })


@login_required
@require_GET
@never_cache
def view_uploaded_disposisi(request, pk):
    disposisi = _uploaded_disposisi_or_404(request.user, pk)
    document = disposisi.dokumen_disposisi
    filename = os.path.basename(document.name)
    content_type = (
        mimetypes.guess_type(filename)[0]
        or 'application/octet-stream'
    )
    return FileResponse(
        document.open('rb'),
        as_attachment=False,
        filename=filename,
        content_type=content_type,
    )


@login_required
@require_GET
@never_cache
def download_uploaded_disposisi(request, pk):
    disposisi = _uploaded_disposisi_or_404(request.user, pk)
    document = disposisi.dokumen_disposisi
    filename = os.path.basename(document.name)
    content_type = (
        mimetypes.guess_type(filename)[0]
        or 'application/octet-stream'
    )
    return FileResponse(
        document.open('rb'),
        as_attachment=True,
        filename=filename,
        content_type=content_type,
    )
