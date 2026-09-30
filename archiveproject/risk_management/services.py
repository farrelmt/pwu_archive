import re
import unicodedata
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .models import (
    RiskActionPlan,
    RiskActivity,
    RiskDivision,
    RiskMonitoring,
    RiskRegister,
    RiskTreatment,
)


REQUIRED_IDENTIFICATION_SHEET = "Identifikasi Risiko"
MONITORING_SHEETS = {
    "Pemantauan TW1": 1,
    "Pemantauan TW2": 2,
    "Pemantauan TW 3": 3,
    "Pemantauan TW3": 3,
    "Pemantauan TW4": 4,
    "Pemantauan TW 4": 4,
}


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _key(value):
    value = unicodedata.normalize("NFKC", _text(value)).casefold()
    return " ".join(value.split())


def _lines(rows, column):
    values = []
    seen = set()
    for row in rows:
        value = _text(_value(row, column))
        if value and value not in seen:
            seen.add(value)
            values.append(value)
    return "\n".join(values)


def _value(row, column):
    return row[column - 1].value if len(row) >= column else None


def _score(value, *, default=None):
    if value in (None, ""):
        return default
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValidationError(f"Nilai skala {value!r} bukan angka 1 sampai 5.")
    if number not in range(1, 6):
        raise ValidationError(f"Nilai skala {number} harus berada di antara 1 dan 5.")
    return number


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(_text(value), pattern).date()
        except ValueError:
            pass
    return None


def _groups(sheet, *, start_row, title_column):
    groups = []
    current = None
    for row_number in range(start_row, sheet.max_row + 1):
        row = list(sheet[row_number])
        title = _text(row[title_column - 1].value)
        number = row[0].value
        if title and (number not in (None, "") or current is None):
            current = {"title": title, "rows": []}
            groups.append(current)
        if current and any(cell.value not in (None, "") for cell in row):
            current["rows"].append(row)
    return groups


def _division_for(name):
    normalized = _key(name).removeprefix("divisi ")
    for division in RiskDivision.objects.filter(is_active=True):
        division_name = _key(division.name).removeprefix("divisi ")
        if normalized in {division_name, _key(division.code)}:
            return division
    raise ValidationError(
        f"Divisi {name!r} pada workbook belum terdaftar di sistem."
    )


def _category_for(title, impacts):
    content = _key(f"{title} {impacts}")
    rules = (
        ("financial", ("keuangan", "finansial", "kerugian", "rupiah")),
        ("compliance", ("hukum", "regulasi", "kepatuhan", "sop")),
        ("reputation", ("reputasi", "nama baik", "citra")),
        ("technology", ("teknologi", "sistem informasi", "data", "siber")),
        ("strategic", ("strategis", "sasaran perusahaan")),
    )
    for category, keywords in rules:
        if any(keyword in content for keyword in keywords):
            return category
    return "operational"


def _marked(value):
    return bool(_text(value))


def _year_from_filename(filename):
    matches = re.findall(r"\b(20\d{2})\b", filename or "")
    return int(matches[-1]) if matches else date.today().year


def _realization_state(realization):
    content = _key(realization)
    if not content or any(word in content for word in ("belum", "tidak dilakukan")):
        return 0, "planned"
    if any(word in content for word in ("telah", "selesai", "dilaksanakan")):
        return 100, "completed"
    return 50, "in_progress"


@transaction.atomic
def import_risk_workbook(uploaded_file, *, user, allowed_division_ids=None):
    try:
        workbook = load_workbook(uploaded_file, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationError(f"Workbook tidak dapat dibaca: {exc}") from exc

    if REQUIRED_IDENTIFICATION_SHEET not in workbook.sheetnames:
        raise ValidationError(
            f"Sheet {REQUIRED_IDENTIFICATION_SHEET!r} tidak ditemukan."
        )

    identification = workbook[REQUIRED_IDENTIFICATION_SHEET]
    division = _division_for(identification["C4"].value)
    if allowed_division_ids is not None and division.pk not in allowed_division_ids:
        raise ValidationError("Anda tidak memiliki akses impor untuk divisi pada workbook ini.")

    owner = _text(identification["C5"].value)
    officer = _text(identification["C6"].value)
    analysis_groups = {
        _key(group["title"]): group
        for group in _groups(
            workbook["Analisis & Evaluasi Risiko"], start_row=10, title_column=2
        )
    } if "Analisis & Evaluasi Risiko" in workbook.sheetnames else {}
    treatment_groups = {
        _key(group["title"]): group
        for group in _groups(workbook["Perlakuan Risiko"], start_row=10, title_column=2)
    } if "Perlakuan Risiko" in workbook.sheetnames else {}

    stats = {
        "risks_created": 0,
        "risks_updated": 0,
        "treatments_created": 0,
        "monitorings_created": 0,
        "actions_created": 0,
    }
    risks_by_title = {}

    for group in _groups(identification, start_row=10, title_column=3):
        rows = group["rows"]
        first = rows[0]
        title = group["title"]
        analysis = analysis_groups.get(_key(title), {"rows": []})["rows"]
        objectives = _lines(rows, 2)
        indications = _lines(rows, 4)
        causes = _lines(rows, 5)
        impacts = _lines(rows, 7)
        controls = _lines(analysis, 3)
        effectiveness = ""
        if any(_marked(row[5].value) for row in analysis):
            effectiveness = "effective"
        elif any(_marked(row[4].value) for row in analysis):
            effectiveness = "adequate"
        elif any(_marked(row[3].value) for row in analysis):
            effectiveness = "ineffective"

        existing = RiskRegister.objects.filter(
            division=division, title__iexact=title
        ).first()
        risk = existing or RiskRegister(division=division, title=title, created_by=user)
        risk.category = existing.category if existing else _category_for(title, impacts)
        risk.strategic_objective = objectives
        risk.description = indications or title
        risk.indication = indications
        risk.cause = causes or "-"
        controllability = _key(_lines(rows, 6))
        risk.controllability = (
            "uncontrollable" if "uncontrollable" in controllability
            else "controllable" if "controllable" in controllability else ""
        )
        risk.impact = impacts or "-"
        risk.inherent_likelihood = _score(first[7].value, default=1)
        risk.inherent_impact = _score(first[8].value, default=1)
        risk.existing_controls = controls
        risk.control_effectiveness = effectiveness
        risk.is_priority = any(_marked(row[10].value) for row in analysis)
        risk.risk_owner = owner or "-"
        risk.risk_officer = officer
        analysis_first = analysis[0] if analysis else None
        risk.residual_likelihood = _score(
            analysis_first[6].value if analysis_first else None,
            default=risk.inherent_likelihood,
        )
        risk.residual_impact = _score(
            analysis_first[7].value if analysis_first else None,
            default=risk.inherent_impact,
        )
        if not existing:
            risk.status = "mitigating" if _key(title) in treatment_groups else "open"
        risk.updated_by = user
        risk.save()
        risks_by_title[_key(title)] = risk
        stats["risks_updated" if existing else "risks_created"] += 1

        treatment_rows = treatment_groups.get(_key(title), {"rows": []})["rows"]
        treatment_descriptions = []
        for row in treatment_rows:
            action_plan = _text(row[3].value)
            option = _text(row[2].value)
            if not action_plan and not option:
                continue
            action_plan = action_plan or option
            treatment_descriptions.append(action_plan)
            treatment = RiskTreatment.objects.filter(
                risk=risk, option=option, action_plan=action_plan
            ).first()
            if not treatment:
                treatment = RiskTreatment(
                    risk=risk, option=option, action_plan=action_plan, created_by=user
                )
                stats["treatments_created"] += 1
            treatment.expected_likelihood = _score(_value(row, 5))
            treatment.expected_impact = _score(_value(row, 6))
            treatment.target_date = _date(_value(row, 9))
            treatment.responsible_person = _text(_value(row, 10))
            treatment.notes = _text(_value(row, 11))
            treatment.updated_by = user
            treatment.save()
        if treatment_descriptions:
            risk.mitigation_plan = "\n".join(treatment_descriptions)
            risk.target_date = min(
                (item.target_date for item in risk.treatments.all() if item.target_date),
                default=None,
            )
            risk.save(update_fields=["mitigation_plan", "target_date", "updated_at"])

        RiskActivity.objects.create(
            risk=risk,
            actor=user,
            action="workbook_imported",
            description="Data risiko diselaraskan dari workbook Manajemen Risiko.",
        )

    year = _year_from_filename(getattr(uploaded_file, "name", ""))
    for sheet_name, quarter in MONITORING_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for group in _groups(sheet, start_row=10, title_column=2):
            risk = risks_by_title.get(_key(group["title"]))
            if not risk:
                continue
            rows = group["rows"]
            first = rows[0]
            defaults = {
                "initial_likelihood": _score(first[6].value, default=risk.residual_likelihood),
                "initial_impact": _score(first[7].value, default=risk.residual_impact),
                "final_likelihood": _score(first[10].value, default=risk.residual_likelihood),
                "final_impact": _score(first[11].value, default=risk.residual_impact),
                "expected_likelihood": _score(first[14].value),
                "expected_impact": _score(first[15].value),
                "evaluation_notes": _lines(rows, 20),
                "updated_by": user,
            }
            monitoring, created = RiskMonitoring.objects.update_or_create(
                risk=risk, year=year, quarter=quarter, defaults=defaults
            )
            stats["monitorings_created"] += int(created)
            for row in rows:
                description = _text(row[2].value)
                if not description:
                    continue
                action = RiskActionPlan.objects.filter(
                    monitoring=monitoring, description=description
                ).first()
                if not action:
                    action = RiskActionPlan(
                        monitoring=monitoring, description=description, created_by=user
                    )
                    stats["actions_created"] += 1
                realization = _text(row[3].value)
                progress, status = _realization_state(realization)
                action.responsible_person = _text(row[5].value) or risk.risk_owner
                action.target_date = _date(row[4].value)
                action.realization = realization
                action.progress = progress
                action.status = status
                action.updated_by = user
                action.save()

    workbook.close()
    stats["division"] = division.name
    return stats
