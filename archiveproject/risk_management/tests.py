from datetime import datetime
from io import BytesIO
from xml.etree import ElementTree

from django.contrib.auth import get_user_model
from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from zipfile import ZipFile

from accounts.models import ActivityLog

from .models import (
    RiskAccess, RiskActionPlan, RiskDivision, RiskMonitoring, RiskRegister,
    RiskTreatment,
)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
}, ROOT_URLCONF="risk_management.host_urls")
class RiskPermissionTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.division_a = RiskDivision.objects.create(code="A", name="Divisi A")
        self.division_b = RiskDivision.objects.create(code="B", name="Divisi B")
        self.officer = User.objects.create_user(username="officer", password="test-pass-123", role="risk")
        self.head = User.objects.create_user(username="head", password="test-pass-123", role="risk")
        RiskAccess.objects.create(user=self.officer, role="risk_officer", division=self.division_a)
        RiskAccess.objects.create(user=self.head, role="division_head", division=self.division_a)
        self.risk = RiskRegister.objects.create(
            division=self.division_a, title="Gangguan operasional", category="operational",
            description="Uraian", cause="Penyebab", impact="Dampak",
            inherent_likelihood=4, inherent_impact=4, existing_controls="Kontrol",
            mitigation_plan="Mitigasi", risk_owner="Pemilik", residual_likelihood=2,
            residual_impact=3, status="mitigating", created_by=self.officer, updated_by=self.officer,
        )

    def test_officer_can_edit_own_division(self):
        self.client.force_login(self.officer)
        response = self.client.get(reverse("risk:risk_edit", args=[self.risk.pk]), HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 200)

    def test_managers_can_replace_division_pic_and_audit_access_changes(self):
        User = get_user_model()
        replacement = User.objects.create_user(username="new-pic", password="test-pass-123", role="risk")
        other_access = RiskAccess.objects.create(user=self.officer, role="risk_officer", division=self.division_b)
        for role in ("head_manager", "manager_member"):
            manager = User.objects.create_user(username=role, password="test-pass-123", role="risk")
            manager_access = RiskAccess.objects.create(user=manager, role=role)
            self.client.force_login(manager)
            listing = self.client.get(reverse("risk:divisions"), HTTP_HOST="risk.localhost")
            self.assertContains(listing, "Edit PIC")
            url = reverse("risk:division_pic_edit", args=[self.division_a.pk])
            self.assertEqual(self.client.get(url, HTTP_HOST="risk.localhost").status_code, 200)
            response = self.client.post(url, {"risk_officers": [replacement.pk], "division_heads": [self.head.pk]}, HTTP_HOST="risk.localhost")
            self.assertRedirects(response, reverse("risk:divisions"), fetch_redirect_response=False)
            self.assertTrue(RiskAccess.objects.filter(user=replacement, division=self.division_a, role="risk_officer", is_active=True).exists())
            self.assertFalse(RiskAccess.objects.get(user=self.officer, division=self.division_a, role="risk_officer").is_active)
            other_access.refresh_from_db()
            manager_access.refresh_from_db()
            self.assertTrue(other_access.is_active)
            self.assertTrue(manager_access.is_active)
        self.assertEqual(ActivityLog.objects.filter(action="DIVISION_PIC_UPDATED").count(), 1)
        log = ActivityLog.objects.get(action="DIVISION_PIC_UPDATED")
        self.assertEqual(log.metadata["before"]["risk_officer"], [self.officer.pk])
        self.assertEqual(log.metadata["after"]["risk_officer"], [replacement.pk])
        self.client.force_login(self.officer)
        self.assertEqual(self.client.get(reverse("risk:risk_detail", args=[self.risk.pk]), HTTP_HOST="risk.localhost").status_code, 404)
        self.client.force_login(replacement)
        self.assertEqual(self.client.get(reverse("risk:risk_edit", args=[self.risk.pk]), HTTP_HOST="risk.localhost").status_code, 200)

    def test_pic_edit_rejects_non_managers_and_inactive_users(self):
        url = reverse("risk:division_pic_edit", args=[self.division_a.pk])
        for user in (self.officer, self.head):
            self.client.force_login(user)
            self.assertNotContains(self.client.get(reverse("risk:divisions"), HTTP_HOST="risk.localhost"), "Edit PIC")
            self.assertEqual(self.client.get(url, HTTP_HOST="risk.localhost").status_code, 403)
            self.assertEqual(self.client.post(url, {}, HTTP_HOST="risk.localhost").status_code, 403)
        admin = get_user_model().objects.create_superuser(username="pic-admin", password="test-pass-123")
        inactive = get_user_model().objects.create_user(username="inactive-pic", is_active=False, role="risk")
        self.client.force_login(admin)
        response = self.client.post(url, {"risk_officers": [inactive.pk]}, HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].errors)
        self.assertTrue(RiskAccess.objects.get(user=self.officer, division=self.division_a, role="risk_officer").is_active)
        self.assertFalse(ActivityLog.objects.filter(action="DIVISION_PIC_UPDATED").exists())
        self.client.post(url, {}, HTTP_HOST="risk.localhost")
        self.assertFalse(RiskAccess.objects.filter(division=self.division_a, is_active=True).exists())
        self.client.post(url, {"risk_officers": [self.officer.pk]}, HTTP_HOST="risk.localhost")
        self.assertEqual(RiskAccess.objects.filter(user=self.officer, division=self.division_a, role="risk_officer").count(), 1)
        self.assertTrue(RiskAccess.objects.get(user=self.officer, division=self.division_a, role="risk_officer").is_active)

    def test_document_form_creates_risk_without_extra_fields(self):
        self.client.force_login(self.officer)
        payload = {
            "division": self.division_a.pk, "title": "Risiko sesuai dokumen",
            "strategic_objective": "Sasaran", "indication": "Indikasi",
            "cause": "Sebab", "controllability": "controllable", "impact": "Dampak",
            "inherent_likelihood": 4, "inherent_impact": 4,
            "existing_controls": "Kontrol", "control_effectiveness": "adequate",
            "residual_likelihood": 2, "residual_impact": 3, "is_priority": "on",
            "risk_owner": "Pemilik", "risk_officer": "Officer",
            "identification_notes": "Catatan identifikasi", "evaluation_notes": "Catatan evaluasi",
        }
        response = self.client.post(reverse("risk:risk_create"), payload, HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 302)
        risk = RiskRegister.objects.get(title=payload["title"])
        self.assertEqual(risk.identification_notes, payload["identification_notes"])
        self.assertEqual(risk.evaluation_notes, payload["evaluation_notes"])
        self.assertEqual(risk.description, payload["indication"])
        self.assertEqual(risk.status, "open")

        risk.category = "financial"
        risk.status = "closed"
        risk.mitigation_plan = "Mitigasi sebelumnya"
        risk.description = "Uraian sebelumnya"
        risk.save()
        payload["indication"] = "Indikasi diperbarui"
        response = self.client.post(reverse("risk:risk_edit", args=[risk.pk]), payload, HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 302)
        risk.refresh_from_db()
        self.assertEqual(risk.category, "financial")
        self.assertEqual(risk.status, "closed")
        self.assertEqual(risk.mitigation_plan, "Mitigasi sebelumnya")
        self.assertEqual(risk.description, "Uraian sebelumnya")
        self.assertEqual(risk.indication, "Indikasi diperbarui")

    def test_document_form_only_renders_document_fields(self):
        self.client.force_login(self.officer)
        response = self.client.get(reverse("risk:risk_create"), HTTP_HOST="risk.localhost")
        self.assertContains(response, "Analisis &amp; Evaluasi Risiko")
        self.assertContains(response, "Risiko Inheren (risiko sebelum pengendalian)")
        self.assertContains(response, "Risiko Setelah Pengendalian yang Ada")
        for name in ("category", "description", "status", "mitigation_plan", "target_date"):
            self.assertNotContains(response, f'name="{name}"')
        self.assertContains(response, 'name="control_effectiveness"', count=3)

    def test_division_head_is_read_only(self):
        self.client.force_login(self.head)
        response = self.client.get(reverse("risk:risk_edit", args=[self.risk.pk]), HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 403)

    def test_division_head_cannot_see_other_division(self):
        other = RiskRegister.objects.create(
            division=self.division_b, title="Risiko lain", category="financial",
            description="Uraian", cause="Penyebab", impact="Dampak",
            inherent_likelihood=3, inherent_impact=3, risk_owner="Pemilik",
            residual_likelihood=2, residual_impact=2, created_by=self.officer, updated_by=self.officer,
        )
        self.client.force_login(self.head)
        response = self.client.get(reverse("risk:risk_detail", args=[other.pk]), HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 404)

    def test_risk_score_and_level(self):
        self.assertEqual(self.risk.inherent_score, 16)
        self.assertEqual(self.risk.residual_score, 6)
        self.assertEqual(self.risk.risk_level, "Sedang")

    def _annual_download(self, user=None, **params):
        self.client.force_login(user or self.officer)
        return self.client.get(reverse("risk:annual_report"), {
            "year": self.risk.created_at.year, "division": self.division_a.pk,
            "download": "1", **params,
        }, HTTP_HOST="risk.localhost")

    def test_annual_excel_preserves_template_and_exports_selected_year(self):
        from .exports import TEMPLATE
        year = self.risk.created_at.year
        self.risk.title = "=1+1"
        self.risk.save()
        monitoring = RiskMonitoring.objects.create(
            risk=self.risk, year=year, quarter=1, initial_likelihood=4, initial_impact=4,
            final_likelihood=2, final_impact=2, expected_likelihood=2, expected_impact=2,
            updated_by=self.officer,
        )
        RiskActionPlan.objects.create(
            monitoring=monitoring, description="Rencana tahun terpilih",
            responsible_person="Pemilik", target_date=f"{year}-03-31",
            created_by=self.officer, updated_by=self.officer,
        )
        RiskActionPlan.objects.create(
            monitoring=monitoring, description="Aksi lanjutan",
            responsible_person="Pemilik", target_date=f"{year}-06-30",
            created_by=self.officer, updated_by=self.officer,
        )
        other_period = RiskMonitoring.objects.create(
            risk=self.risk, year=year - 1, quarter=1, initial_likelihood=1, initial_impact=1,
            final_likelihood=1, final_impact=1, updated_by=self.officer,
        )
        RiskActionPlan.objects.create(
            monitoring=other_period, description="Jangan ekspor tahun lain",
            responsible_person="Pemilik", created_by=self.officer, updated_by=self.officer,
        )
        response = self._annual_download()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        payload = b"".join(response.streaming_content)
        exported = load_workbook(BytesIO(payload), data_only=False)
        values = load_workbook(BytesIO(payload), data_only=True)
        original = load_workbook(TEMPLATE)
        self.assertEqual(exported.sheetnames, original.sheetnames)
        self.assertEqual(exported["Identifikasi Risiko"]["C10"].value, "=1+1")
        self.assertEqual(exported["Identifikasi Risiko"]["C10"].data_type, "s")
        self.assertEqual(values["Identifikasi Risiko"]["J10"].value, 16)
        self.assertEqual(values["Pemantauan TW1"]["S10"].value, 0)
        self.assertEqual(values["Pemantauan TW1"]["C10"].value, "Rencana tahun terpilih")
        self.assertEqual(values["Pemantauan TW1"]["E10"].value.date().isoformat(), f"{year}-03-31")
        self.assertEqual(values["Pemantauan TW1"]["E11"].value.date().isoformat(), f"{year}-06-30")
        self.assertEqual(values["Identifikasi Risiko"]["C4"].value, self.division_a.name)
        for name in exported.sheetnames[3:]:
            self.assertEqual(exported[name]["A7"].value, original[name]["A7"].value)
            self.assertEqual(exported[name].column_dimensions["B"].width, original[name].column_dimensions["B"].width)
        with ZipFile(BytesIO(payload)) as actual, ZipFile(TEMPLATE) as source:
            for name in ["xl/styles.xml", "xl/theme/theme1.xml", "xl/drawings/drawing1.xml", *[f"xl/worksheets/sheet{i}.xml" for i in (1, 2, 3)]]:
                self.assertEqual(actual.read(name), source.read(name))
        self.assertNotIn("Jangan ekspor tahun lain", str(list(values["Pemantauan TW1"].values)))
        exported.close()
        values.close()
        original.close()

    def test_annual_excel_restricts_division_and_validates_year(self):
        response = self._annual_download(division=self.division_b.pk)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.streaming)
        self.assertTrue(response.context["form"].errors)
        response = self._annual_download(year="invalid")
        self.assertTrue(response.context["form"].errors)
        response = self._annual_download(year=2020)
        self.assertContains(response, "Tidak ada data risiko")
        self.assertEqual(ActivityLog.objects.filter(
            actor=self.officer, action="ANNUAL_REPORT_DOWNLOAD_FAILED", success=False,
        ).count(), 3)
        self.assertFalse(ActivityLog.objects.filter(action="ANNUAL_REPORT_DOWNLOADED").exists())

    def test_download_audit_records_actor_document_and_request_once(self):
        self.client.force_login(self.officer)
        url = reverse("risk:annual_report")
        self.client.get(url, HTTP_HOST="risk.localhost")
        self.assertFalse(ActivityLog.objects.filter(category="RISK").exists())
        started = timezone.now()
        response = self.client.get(url, {
            "year": self.risk.created_at.year, "division": self.division_a.pk, "download": "1",
        }, HTTP_HOST="risk.localhost", REMOTE_ADDR="192.0.2.10", HTTP_USER_AGENT="Audit test browser")
        self.assertTrue(response.streaming)
        log = ActivityLog.objects.get(category="RISK")
        self.assertEqual(log.action, "ANNUAL_REPORT_DOWNLOADED")
        self.assertEqual(log.actor, self.officer)
        self.assertEqual(log.actor_username, "officer")
        self.assertGreaterEqual(log.created_at, started)
        self.assertEqual(log.ip_address, "192.0.2.10")
        self.assertEqual(log.user_agent, "Audit test browser")
        self.assertEqual(log.metadata["year"], self.risk.created_at.year)
        self.assertEqual(log.metadata["division_id"], self.division_a.pk)
        self.assertEqual(log.metadata["risk_count"], 1)
        self.assertEqual(log.metadata["delivery"], "response_prepared")
        self.assertEqual(log.metadata["path"], url)
        self.assertIn(log.metadata["filename"], response["Content-Disposition"])
        self.assertEqual(log.target_label, log.metadata["filename"])
        self.assertTrue(log.success)
        response.close()

    def test_view_audit_records_visible_risk_and_report_filters(self):
        self.client.force_login(self.officer)
        response = self.client.get(reverse("risk:risk_detail", args=[self.risk.pk]), HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 200)
        log = ActivityLog.objects.get(action="RISK_VIEWED")
        self.assertEqual(log.target_id, str(self.risk.pk))
        self.assertEqual(log.target_label, self.risk.risk_code)
        self.client.get(reverse("risk:risk_list"), {"year": self.risk.created_at.year}, HTTP_HOST="risk.localhost")
        self.assertEqual(ActivityLog.objects.get(action="RISK_REGISTER_VIEWED").metadata["year"], str(self.risk.created_at.year))
        self.client.get(reverse("risk:monitoring_report"), {"year": 2026, "quarter": 2}, HTTP_HOST="risk.localhost")
        log = ActivityLog.objects.get(action="MONITORING_REPORT_VIEWED")
        self.assertEqual(log.metadata["year"], 2026)
        self.assertEqual(log.metadata["quarter"], 2)
        other = RiskRegister.objects.create(
            division=self.division_b, title="Rahasia", category="operational",
            inherent_likelihood=1, inherent_impact=1, residual_likelihood=1, residual_impact=1,
            created_by=self.officer, updated_by=self.officer,
        )
        response = self.client.get(reverse("risk:risk_detail", args=[other.pk]), HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 404)
        self.assertEqual(ActivityLog.objects.filter(action="RISK_VIEWED").count(), 1)

    def test_download_mark_identifies_each_recipient_without_changing_sheets(self):
        from .exports import DOWNLOAD_MARK_SALT, annual_risk_workbook

        self.officer.first_name = "Rina & <PWU>"
        self.officer.save()
        baseline = annual_risk_workbook([self.risk], division=self.division_a, year=self.risk.created_at.year)
        marks = []
        for user in (self.officer, self.officer, self.head):
            response = self._annual_download(user=user)
            payload = b"".join(response.streaming_content)
            response.close()
            book = load_workbook(BytesIO(payload))
            properties = {prop.name: prop.value for prop in book.custom_doc_props}
            book.close()
            mark = signing.loads(properties["RiskDownloadProvenance"], salt=DOWNLOAD_MARK_SALT)
            marks.append(mark)
            self.assertEqual(mark["username"], user.username)
            self.assertEqual(mark["user_id"], str(user.pk))
            self.assertEqual(mark["full_name"], user.get_full_name() or user.username)
            self.assertEqual(properties["RiskDownloadedBy"], mark["username"])
            self.assertEqual(properties["RiskDownloaderName"], mark["full_name"])
            self.assertEqual(properties["RiskDownloadedAt"], mark["downloaded_at"])
            self.assertEqual(datetime.fromisoformat(mark["downloaded_at"]).utcoffset().total_seconds(), 7 * 3600)
            log = ActivityLog.objects.get(action="ANNUAL_REPORT_DOWNLOADED", metadata__download_mark__download_id=mark["download_id"])
            self.assertEqual(log.actor, user)
            self.assertEqual(log.metadata["download_mark"], mark)
            token = properties["RiskDownloadProvenance"]
            with self.assertRaises(signing.BadSignature):
                signing.loads(token + "tampered", salt=DOWNLOAD_MARK_SALT)
            with ZipFile(BytesIO(payload)) as marked, ZipFile(baseline) as original:
                for name in original.namelist():
                    if name not in {"docProps/custom.xml", "_rels/.rels", "[Content_Types].xml"}:
                        self.assertEqual(marked.read(name), original.read(name), name)
                # Excel must be able to discover the custom properties part.
                rels = ElementTree.fromstring(marked.read("_rels/.rels"))
                self.assertEqual(sum(rel.attrib["Type"].endswith("/custom-properties") for rel in rels), 1)
                types = ElementTree.fromstring(marked.read("[Content_Types].xml"))
                self.assertEqual(sum(node.attrib.get("PartName") == "/docProps/custom.xml" for node in types), 1)
        self.assertEqual(len({mark["download_id"] for mark in marks}), 3)
        baseline.close()

    def test_annual_excel_adds_fourth_quarter_without_sample_rows(self):
        RiskMonitoring.objects.create(
            risk=self.risk, year=self.risk.created_at.year, quarter=4,
            initial_likelihood=4, initial_impact=4, final_likelihood=2, final_impact=2,
            updated_by=self.officer,
        )
        response = self._annual_download(user=self.head)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)), data_only=True)
        self.assertIn("Pemantauan TW4", workbook.sheetnames)
        self.assertEqual(workbook["Pemantauan TW4"]["B10"].value, self.risk.title)
        self.assertIsNone(workbook["Pemantauan TW4"]["B16"].value)
        self.assertEqual(workbook["Pemantauan TW4"]["M10"].value, 4)
        workbook.close()

    def test_register_sorting_applies_before_pagination_and_keeps_filters(self):
        for index in range(21):
            RiskRegister.objects.create(
                division=self.division_a, title=f"Risiko {index:02}", category="operational",
                description="Uraian", cause="Sebab", impact="Dampak", risk_owner="Pemilik",
                inherent_likelihood=1, inherent_impact=1,
                residual_likelihood=5 if index == 20 else 1, residual_impact=5,
                created_by=self.officer, updated_by=self.officer,
            )
        self.client.force_login(self.officer)
        response = self.client.get(reverse("risk:risk_list"), {
            "sort": "residual", "order": "desc", "q": "Risiko", "division": self.division_a.pk,
        }, HTTP_HOST="risk.localhost")
        page = response.context["page_obj"]
        self.assertEqual(page[0].title, "Risiko 20")
        self.assertEqual(page.paginator.count, 21)
        self.assertContains(response, "sort=residual&amp;order=desc")
        self.assertIn("division=", response.context["pagination_query"])
        ascending = self.client.get(reverse("risk:risk_list"), {
            "sort": "residual", "order": "asc", "q": "Risiko",
        }, HTTP_HOST="risk.localhost")
        self.assertEqual(ascending.context["page_obj"][0].residual_score, 5)

    def test_monitoring_sorting_keeps_actions_grouped_and_missing_targets_last(self):
        other = RiskRegister.objects.create(
            division=self.division_a, title="Autre risque", category="operational",
            description="Uraian", cause="Sebab", impact="Dampak", risk_owner="Pemilik",
            inherent_likelihood=1, inherent_impact=1, residual_likelihood=1, residual_impact=1,
            created_by=self.officer, updated_by=self.officer,
        )
        missing = RiskMonitoring.objects.create(
            risk=other, year=2026, quarter=2, initial_likelihood=1, initial_impact=1,
            final_likelihood=1, final_impact=1, updated_by=self.officer,
        )
        monitoring = RiskMonitoring.objects.create(
            risk=self.risk, year=2026, quarter=2, initial_likelihood=4, initial_impact=4,
            final_likelihood=3, final_impact=3, expected_likelihood=2, expected_impact=2,
            updated_by=self.officer,
        )
        for description in ("Z aksi", "A aksi"):
            RiskActionPlan.objects.create(
                monitoring=monitoring, description=description, responsible_person="Pemilik",
                created_by=self.officer, updated_by=self.officer,
            )
        self.client.force_login(self.officer)
        for direction in ("asc", "desc"):
            response = self.client.get(reverse("risk:monitoring_report"), {
                "year": 2026, "quarter": 2, "sort": "expected_score", "order": direction,
            }, HTTP_HOST="risk.localhost")
            self.assertEqual([m.pk for m in response.context["monitorings"]], [monitoring.pk, missing.pk])
        response = self.client.get(reverse("risk:monitoring_report"), {
            "year": 2026, "quarter": 2, "sort": "plan", "order": "asc",
        }, HTTP_HOST="risk.localhost")
        self.assertEqual(response.context["monitorings"][0].pk, monitoring.pk)
        self.assertEqual([a.description for a in response.context["monitorings"][0].action_plans.all()], ["A aksi", "Z aksi"])
        self.assertContains(response, "year=2026&amp;quarter=2")

    def test_dashboard_uses_database_summary(self):
        self.client.force_login(self.officer)

        response = self.client.get(
            reverse("risk:dashboard"),
            HTTP_HOST="risk.localhost",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_risks"], 1)
        self.assertEqual(response.context["open_risks"], 1)
        self.assertEqual(response.context["level_counts"]["Sedang"], 1)
        self.assertContains(response, "Kembali ke Portal Utama")
        self.assertContains(response, ">Keluar<")

    def test_template_uses_five_risk_levels(self):
        expected = {
            1: "Sangat Rendah", 2: "Sangat Rendah", 3: "Rendah", 4: "Rendah",
            5: "Sedang", 9: "Sedang", 10: "Tinggi", 16: "Tinggi",
            17: "Sangat Tinggi", 25: "Sangat Tinggi",
        }
        for score, level in expected.items():
            with self.subTest(score=score):
                self.assertEqual(RiskRegister.level_for(score), level)

    def _workbook_upload(self):
        workbook = Workbook()
        identification = workbook.active
        identification.title = "Identifikasi Risiko"
        identification["C4"] = "Divisi A"
        identification["C5"] = "Pemilik A"
        identification["C6"] = "Officer A"
        values = {
            "A10": 1, "B10": "Sasaran strategis", "C10": "Risiko dari workbook",
            "D10": "Indikasi", "E10": "Penyebab", "F10": "Controllable",
            "G10": "Dampak", "H10": 4, "I10": 4,
        }
        for cell, value in values.items():
            identification[cell] = value

        analysis = workbook.create_sheet("Analisis & Evaluasi Risiko")
        values = {
            "A10": 1, "B10": "Risiko dari workbook", "C10": "Kontrol A",
            "E10": "✓", "G10": 3, "H10": 3, "K10": "✓",
        }
        for cell, value in values.items():
            analysis[cell] = value

        treatment = workbook.create_sheet("Perlakuan Risiko")
        values = {
            "A10": 1, "B10": "Risiko dari workbook", "C10": "Mitigasi",
            "D10": "Perbarui SOP", "E10": 2, "F10": 2,
            "I10": "2026-12-31", "J10": "Pemilik A",
        }
        for cell, value in values.items():
            treatment[cell] = value

        monitoring = workbook.create_sheet("Pemantauan TW1")
        values = {
            "A10": 1, "B10": "Risiko dari workbook", "C10": "Perbarui SOP",
            "D10": "Telah dilaksanakan", "E10": "2026-12-31", "F10": "Pemilik A",
            "G10": 3, "H10": 3, "K10": 2, "L10": 2, "O10": 2, "P10": 2,
        }
        for cell, value in values.items():
            monitoring[cell] = value

        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return SimpleUploadedFile(
            "Manajemen Risiko 2026.xlsx",
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_officer_can_import_template_without_duplicates(self):
        self.client.force_login(self.officer)
        for expected_status in (302, 302):
            response = self.client.post(
                reverse("risk:workbook_import"),
                {"workbook": self._workbook_upload()},
                HTTP_HOST="risk.localhost",
            )
            self.assertEqual(response.status_code, expected_status)
        imported = RiskRegister.objects.get(title="Risiko dari workbook")
        self.assertEqual(imported.strategic_objective, "Sasaran strategis")
        self.assertEqual(imported.control_effectiveness, "adequate")
        self.assertTrue(imported.is_priority)
        self.assertEqual(imported.risk_level, "Sedang")
        self.assertEqual(RiskTreatment.objects.filter(risk=imported).count(), 1)
        monitoring = RiskMonitoring.objects.get(risk=imported, year=2026, quarter=1)
        self.assertEqual(monitoring.expected_score, 4)
        self.assertEqual(RiskActionPlan.objects.filter(monitoring=monitoring).count(), 1)
        self.assertEqual(ActivityLog.objects.filter(action="RISK_WORKBOOK_IMPORTED", actor=self.officer, success=True).count(), 2)
        self.assertEqual(ActivityLog.objects.filter(action="RISK_WORKBOOK_IMPORTED").first().metadata["filename"], "Manajemen Risiko 2026.xlsx")

    def test_import_rejects_other_division(self):
        source = Workbook()
        source.active.title = "Identifikasi Risiko"
        source.active["C4"] = "Divisi B"
        stream = BytesIO()
        source.save(stream)
        source.close()
        denied_workbook = SimpleUploadedFile("Risiko 2026.xlsx", stream.getvalue())
        self.client.force_login(self.officer)
        response = self.client.post(
            reverse("risk:workbook_import"), {"workbook": denied_workbook},
            HTTP_HOST="risk.localhost",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "tidak memiliki akses impor")
        self.assertTrue(ActivityLog.objects.filter(action="RISK_WORKBOOK_IMPORT_FAILED", actor=self.officer, success=False).exists())

    def test_officer_can_record_quarterly_monitoring(self):
        self.client.force_login(self.officer)
        response = self.client.post(
            reverse("risk:monitoring_create", args=[self.risk.pk]),
            {
                "year": 2026, "quarter": 2,
                "initial_likelihood": 4, "initial_impact": 4,
                "final_likelihood": 3, "final_impact": 3,
                "business_environment_changes": "Tidak ada perubahan material.",
                "evaluation_notes": "Mitigasi berjalan.",
            },
            HTTP_HOST="risk.localhost",
        )
        self.assertEqual(response.status_code, 302)
        monitoring = RiskMonitoring.objects.get(risk=self.risk, year=2026, quarter=2)
        self.assertEqual(monitoring.initial_score, 16)
        self.assertEqual(monitoring.final_score, 9)
        self.assertEqual(monitoring.score_reduction, 7)
        self.assertTrue(ActivityLog.objects.filter(
            actor=self.officer,
            category="RISK",
            action="MONITORING_CREATED",
            target_id=str(monitoring.pk),
        ).exists())

    def test_only_superuser_can_open_risk_activity_log(self):
        self.client.force_login(self.officer)
        denied = self.client.get(
            reverse("risk:activity_log"),
            HTTP_HOST="risk.localhost",
        )
        self.assertEqual(denied.status_code, 403)

        admin = get_user_model().objects.create_superuser(
            username="risk-log-admin",
            password="test-pass-123",
        )
        self.client.force_login(admin)
        response = self.client.get(
            reverse("risk:activity_log"),
            HTTP_HOST="risk.localhost",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Log Aktivitas")

    def test_activity_log_filters_downloads_and_displays_jakarta_time(self):
        response = self._annual_download()
        response.close()
        log = ActivityLog.objects.get(action="ANNUAL_REPORT_DOWNLOADED")
        ActivityLog.objects.create(
            actor=self.head, actor_username=self.head.username, category="RISK",
            action="RISK_UPDATED", description="Another action",
        )
        admin = get_user_model().objects.create_superuser(username="audit-admin", password="test-pass-123")
        self.client.force_login(admin)
        today = timezone.localtime(log.created_at).date().isoformat()
        response = self.client.get(reverse("risk:activity_log"), {
            "actor": self.officer.pk, "action": log.action, "result": "success",
            "date_from": today, "date_to": today,
        }, HTTP_HOST="risk.localhost")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context["page_obj"]), [log])
        self.assertContains(response, "Waktu (WIB)")
        self.assertContains(response, timezone.localtime(log.created_at).strftime("%d/%m/%Y %H:%M:%S"))
        self.assertContains(response, log.target_label)
        self.assertContains(response, log.description)

    def test_monitoring_report_summarizes_actions_and_scores(self):
        monitoring = RiskMonitoring.objects.create(
            risk=self.risk, year=2026, quarter=2,
            initial_likelihood=4, initial_impact=4,
            final_likelihood=3, final_impact=3,
            updated_by=self.officer,
        )
        RiskActionPlan.objects.create(
            monitoring=monitoring,
            description="Perbarui prosedur pengendalian.",
            responsible_person="Pemilik",
            realization="Prosedur telah diperbarui.",
            realization_date="2026-06-30",
            progress=100,
            status="completed",
            created_by=self.officer,
            updated_by=self.officer,
        )
        self.client.force_login(self.officer)
        response = self.client.get(
            reverse("risk:monitoring_report"),
            {"year": 2026, "quarter": 2},
            HTTP_HOST="risk.localhost",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["risk_count"], 1)
        self.assertEqual(response.context["totals"]["plan_count"], 1)
        self.assertEqual(response.context["totals"]["realized_count"], 1)
        self.assertEqual(response.context["totals"]["realization_percent"], 100)
        self.assertEqual(response.context["totals"]["score_reduction"], 7)
